import os
import json
import time
import threading
import tempfile
import subprocess
import base64
import re
import requests
from dotenv import load_dotenv
load_dotenv(override=True)
from datetime import datetime, timedelta, timezone
from typing import List, Any, Dict

from langchain_openai import ChatOpenAI
from langchain_core.messages import HumanMessage, AIMessage, SystemMessage

from colors import GREEN, RED, YELLOW, BLUE, RESET
from agent.calendar import criar_evento_google_meet, deletar_evento_google_calendar, verificar_conflito_google_calendar, buscar_horarios_livres, get_organizer_for_product, MASTER_EMAIL
from agent.normalizers import normalizar_data, normalizar_horario, normalizar_telefone
from agent.context import (
    get_contexto,
    PALAVRAS_AGENDAMENTO,
    get_instrucao_followup,
    get_instrucao_abertura,
    get_instrucao_recontato,
    get_instrucao_lembrete,
    get_prompt_extracao_lead,
)
from agent.ana_agent import run_ana_agent
from agent.micro_agents import run_micro_agents, _to_langchain_messages
from agent.storage import supabase_repo
from agent.handoff_opener import gerar_abertura_handoff


FOLLOWUP_OVERDUE_GRACE_SECONDS = 24 * 60 * 60


def _contexto_tem_indicacao_explicita(contexto_lead: str) -> bool:
    """Detect whether lead context explicitly states a referral source.

    Inputs:
        contexto_lead: Free-text lead context used to generate outbound openings.

    Behavior/side effects:
        Normalizes the text in memory and searches for referral terms. Does not
        write files, call APIs or mutate state.

    Return:
        True when the context contains an explicit referral/indication phrase;
        otherwise False.

    Exceptions/fallbacks:
        Empty or None context is treated as no referral. No expected exceptions.
    """
    contexto = (contexto_lead or "").lower()
    termos_indicacao = (
        "indicado por",
        "indicada por",
        "indicacao de",
        "indicação de",
        "quem indicou",
        "me passou seu contato",
        "passou seu contato",
        "foi indicado",
        "foi indicada",
    )
    return any(termo in contexto for termo in termos_indicacao)


def _sanitizar_origem_abertura(mensagem: str, contexto_lead: str) -> str:
    """Remove awkward or invented origin wording from cold outbound openings.

    Inputs:
        mensagem: LLM-generated outbound opening text.
        contexto_lead: Lead context used to generate the opening.

    Behavior/side effects:
        Rewrites only the in-memory message text. Referral-name rewrites are
        skipped when the context explicitly contains an indication source.
        Blocks phrasing that explains how the contact was found.

    Return:
        Sanitized opening text ready to send.

    Exceptions/fallbacks:
        If no banned wording is found, returns the original message unchanged.
        No expected exceptions.
    """
    if _contexto_tem_indicacao_explicita(contexto_lead):
        return mensagem

    origem_neutra = "Tenho falado com operações que buscam reduzir trabalho manual e centralizar rotinas críticas."
    patterns = (
        (r"\b[Mm]eu nome [ée?] Ana e sou da Btime\.?", "Sou a Ana, da Btime."),
        (r"\b[Mm]eu nome [ée?] Ana\.?", "Sou a Ana, da Btime."),
        r"\b(?:O|A)\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ][\wÀ-ÿ.-]*\s+me passou seu contato\.?",
        r"\b[A-ZÁÉÍÓÚÂÊÔÃÕÇ][\wÀ-ÿ.-]*\s+me passou seu contato\.?",
        r"\b(?:O|A)\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ][\wÀ-ÿ.-]*\s+indicou seu contato\.?",
        r"\b[A-ZÁÉÍÓÚÂÊÔÃÕÇ][\wÀ-ÿ.-]*\s+indicou seu contato\.?",
        r"\b[Vv]i seu contexto(?: profissional)?(?:\s+(?:em|como|na|no)\s+[^.\n]+)?\.?",
        r"\b[Oo] sistema trouxe seu contato pelo seu contexto profissional\.?",
        r"\b[Pp]eguei seu contato\b[^.\n]*\.?",
        r"\b[Cc]onsegui seu n[úu]mero\b[^.\n]*\.?",
    )
    sanitized = mensagem
    for pattern in patterns:
        if isinstance(pattern, tuple):
            sanitized = re.sub(pattern[0], pattern[1], sanitized)
            continue
        sanitized = re.sub(pattern, origem_neutra, sanitized)
    return sanitized


def _proximo_horario_util(dt: datetime) -> datetime:
    """Retorna o próximo datetime dentro do horário útil (9-12, 13-19, seg-sex).
    Se dt já estiver dentro do horário útil, retorna dt inalterado."""
    candidate = dt.replace(second=0, microsecond=0)
    for _ in range(10):  # máximo de iterações para evitar loop infinito
        # Fim de semana → avança para segunda-feira
        if candidate.weekday() >= 5:
            dias = 7 - candidate.weekday()
            candidate = datetime(candidate.year, candidate.month, candidate.day, 9, 0) + timedelta(days=dias)
            continue
        hora = candidate.hour + candidate.minute / 60
        if 9 <= hora < 12:
            break
        if 12 <= hora < 13:
            candidate = candidate.replace(hour=13, minute=0)
            break
        if 13 <= hora < 19:
            break
        if hora < 9:
            candidate = candidate.replace(hour=9, minute=0)
            break
        # Após 19h → próximo dia útil às 9h
        candidate = datetime(candidate.year, candidate.month, candidate.day, 9, 0) + timedelta(days=1)
    return candidate


class Agent_AI:

    def __init__(self):
        self.api_token = "E1014DD5B87C05F2DF8F2547"
        self.zapi_sec_token = os.environ.get("ZAPI_SEC_TOKEN")
        self.instance_id = "3EFABF270A13919DE95AD6F4B026E35E"
        self.base_url = f"https://api.z-api.io/instances/{self.instance_id}/token/{self.api_token}"

        self.answer_list = {}
        self.pending_timers: dict[str, threading.Timer] = {}
        self.pending_texts: dict[str, str] = {}
        self.followup_timers: dict[str, list[threading.Timer]] = {}
        self.followup_job_timers: dict[str, threading.Timer] = {}
        self.composing_set: set = set()
        self.processing_chats: set[str] = set()
        self.interrupted_chats: set[str] = set()
        # Numeros com disparo de handoff em voo (chat-scoped). Evita corrida entre
        # turnos enquanto a thread async ainda nao confirmou sucesso. Persistencia
        # final fica na tabela `handoffs` do Supabase (success=true) depois do ok do Z-API.
        self._handoff_in_flight: dict[str, set[str]] = {}
        self._handoff_in_flight_lock = threading.Lock()
        self._interrupt_cite_ids: dict[str, str] = {}
        self._timer_lock = threading.Lock()
        self._file_lock = threading.Lock()
        self.phone_to_lid: dict[str, str] = {}

        self._restaurar_followups()

    def _read_json_file(self, path: str, default):
        with self._file_lock:
            if not os.path.exists(path):
                return default
            try:
                with open(path, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                return default

    def _write_json_file(self, path: str, data):
        with self._file_lock:
            parent = os.path.dirname(path)
            if parent:
                os.makedirs(parent, exist_ok=True)
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2, ensure_ascii=False)

    # ------------------------------------------------------------------
    # Supabase adapters
    #
    # Supabase is the source of truth. Local JSON is kept as a write-mirror
    # (best-effort) and as a read-fallback when Supabase is unavailable.
    # All Supabase calls below MUST go through `_supabase_safe` so that a
    # transient network/HTTP error never crashes the webhook or agent flow.
    # ------------------------------------------------------------------

    def _supabase_safe(self, fn, *args, **kwargs):
        """Call a supabase_repo function and swallow failures.

        Inputs:
            fn: Callable from `supabase_repo`.
            *args, **kwargs: Forwarded to `fn`.

        Returns:
            The function's return value on success, or None on any exception.

        Side effects:
            Logs a short yellow warning on failure. Does NOT log credentials.
        """
        try:
            return fn(*args, **kwargs)
        except Exception as e:
            print(f"{YELLOW}[supabase] {fn.__name__} falhou: {e}{RESET}")
            return None

    def _epoch_from_iso(self, iso_str: str) -> float:
        """Convert an ISO-8601 string to a unix epoch float.

        Inputs:
            iso_str: ISO timestamp (with or without timezone).

        Returns:
            Epoch in seconds. 0.0 when parsing fails.
        """
        if not iso_str:
            return 0.0
        try:
            return datetime.fromisoformat(iso_str).timestamp()
        except Exception:
            return 0.0

    def _supa_load_history(self, chatLid: str) -> list:
        """Read message history for `chatLid` from Supabase, mapped to the legacy shape.

        Inputs:
            chatLid: WhatsApp chat identifier.

        Returns:
            List of `{"timestamp": <epoch>, "data": <z-api payload>}` dicts,
            ordered chronologically. Empty list when nothing is found or
            Supabase is unreachable.
        """
        rows = self._supabase_safe(supabase_repo.load_history, chatLid) or []
        out = []
        for r in rows:
            payload = r.get("payload") or {}
            out.append({
                "timestamp": self._epoch_from_iso(r.get("event_ts") or ""),
                "data": payload,
            })
        return out

    def _supa_load_lead_info(self, chatLid: str):
        """Read lead_info-like dict for `chatLid` from Supabase.

        Inputs:
            chatLid: WhatsApp chat identifier.

        Returns:
            Dict in the legacy lead_info.json shape (merging `lead_profiles`
            columns with `lead_profiles.raw` extras and a couple of fields
            from `conversations`). Returns None if no profile exists yet.
        """
        prof = self._supabase_safe(supabase_repo.get_lead_profile, chatLid)
        if not prof:
            return None
        out: Dict[str, Any] = {}
        for k in (
            "nome", "email", "empresa", "segmento_mercado", "porte_empresa",
            "faturamento_aproximado", "tamanho_time", "sistemas_atuais",
            "desafios_identificados", "produto_indicado", "motivo_segmentacao",
            "estagio_conversa", "reuniao_agendada", "necessita_followup",
            "motivo_followup", "tentativas_retomada", "descricao_manual",
            "atualizado_em",
        ):
            v = prof.get(k)
            if v is not None:
                out[k] = v
        raw = prof.get("raw") or {}
        for k, v in raw.items():
            if k not in out and v is not None:
                out[k] = v
        conv = self._supabase_safe(supabase_repo.get_conversation, chatLid) or {}
        if conv.get("ai_blocked") is not None:
            out["ai_blocked"] = bool(conv["ai_blocked"])
        if conv.get("phone") and "phone" not in out:
            out["phone"] = conv["phone"]
        return out

    def _read_history(self, chatLid: str) -> list:
        """Return history for `chatLid`, preferring Supabase over the JSON file.

        Inputs:
            chatLid: WhatsApp chat identifier.

        Returns:
            List in the legacy `{timestamp, data}` shape. Falls back to
            `chats/{chatLid}/history.json` if Supabase returns nothing.
        """
        rows = self._supa_load_history(chatLid)
        if rows:
            return rows
        return self._read_json_file(os.path.join("chats", chatLid, "history.json"), [])

    def _read_lead_info(self, chatLid: str) -> dict:
        """Return the merged lead_info dict, preferring Supabase.

        Inputs:
            chatLid: WhatsApp chat identifier.

        Returns:
            Dict (possibly empty). Falls back to local `lead_info.json` when
            Supabase has no profile yet.
        """
        info = self._supa_load_lead_info(chatLid)
        if info is not None:
            return info
        return self._read_json_file(os.path.join("chats", chatLid, "lead_info.json"), {})

    def _read_active_meeting_for(self, chatLid: str):
        """Return the next active meeting dict for `chatLid`, or None.

        Inputs:
            chatLid: WhatsApp chat identifier.

        Returns:
            Dict in the legacy `reunioes.json` shape (`chatLid`, `phone`,
            `datetime`, `meet_link`, `event_id`, `agendado_em`) of the nearest
            scheduled meeting for that lead. Falls back to scanning the local
            `reunioes.json` when Supabase has none.
        """
        rows = self._supabase_safe(supabase_repo.list_active_meetings, chatLid) or []
        if rows:
            r = rows[0]
            return {
                "chatLid": chatLid,
                "phone": r.get("phone"),
                "datetime": r.get("scheduled_for"),
                "agendado_em": r.get("agendado_em"),
                "meet_link": r.get("meet_link"),
                "event_id": r.get("event_id"),
            }
        local = self._read_json_file("reunioes.json", [])
        for r in local:
            if r.get("chatLid") == chatLid:
                return r
        return None

    def _has_active_meeting(self, chatLid: str) -> bool:
        try:
            reuniao = self._read_active_meeting_for(chatLid)
            if not reuniao:
                return False
            dt_reuniao = reuniao.get("datetime", "")
            if not dt_reuniao:
                return False
            alvo = datetime.fromisoformat(dt_reuniao)
            if alvo.tzinfo is not None:
                alvo = alvo.astimezone().replace(tzinfo=None)
            return alvo > datetime.now()
        except Exception:
            return False

    def _mirror_message_supabase(self, chatLid: str, event: dict) -> None:
        """Best-effort send of one history event to Supabase `messages`.

        Inputs:
            chatLid: WhatsApp chat identifier.
            event: Dict in the `{timestamp, data}` shape.

        Returns:
            None.

        Side effects:
            Idempotent upsert in Supabase. Never raises; failures are logged
            via `_supabase_safe`.
        """
        if not chatLid or not isinstance(event, dict):
            return
        self._supabase_safe(supabase_repo.append_message, chatLid, event)

    def _mirror_lead_info_supabase(self, chatLid: str, info: dict) -> None:
        """Best-effort sync of a lead_info dict into `conversations` + `lead_profiles`.

        Inputs:
            chatLid: WhatsApp chat identifier.
            info: lead_info-shaped dict.

        Returns:
            None.

        Side effects:
            Upserts the conversation row (carrying phone + ai_blocked) and
            the lead_profile row. `followups_agendados` is stripped: that
            list is managed via `followup_jobs` directly.
        """
        if not chatLid or not isinstance(info, dict):
            return
        self._supabase_safe(
            supabase_repo.upsert_conversation,
            chatLid,
            phone=info.get("phone"),
            ai_blocked=info.get("ai_blocked"),
        )
        profile_data = {
            k: v for k, v in info.items()
            if k not in ("ai_blocked", "phone", "followups_agendados")
        }
        if profile_data:
            self._supabase_safe(supabase_repo.upsert_lead_profile, chatLid, profile_data)

    def _mirror_meeting_supabase(self, chatLid: str, meeting: dict) -> None:
        """Best-effort upsert of a meeting into Supabase `meetings`."""
        if not chatLid or not isinstance(meeting, dict):
            return
        self._supabase_safe(supabase_repo.upsert_meeting, chatLid, meeting)

    def _mirror_meeting_cancelled(self, chatLid: str, iso_dt: str) -> None:
        """Best-effort: mark a meeting as cancelled by (chat_lid, datetime)."""
        if not chatLid or not iso_dt:
            return
        self._supabase_safe(
            supabase_repo.update_meeting_status,
            chat_lid=chatLid,
            datetime=iso_dt,
            status="cancelled",
        )

    def _mirror_followup_scheduled(
        self, chatLid: str, tipo: str, alvo_iso: str, phone: str
    ) -> dict | None:
        """Best-effort: persist a follow-up timer into `followup_jobs`.

        Inputs:
            chatLid: WhatsApp chat identifier.
            tipo: Follow-up type (`1h`, `24h` or `15d`).
            alvo_iso: Target ISO timestamp for the follow-up.
            phone: Destination WhatsApp phone number.

        Side effects:
            Upserts one Supabase `followup_jobs` row with status `scheduled`.

        Returns:
            The Supabase row dict when available, otherwise None.

        Exceptions/fallbacks:
            Never raises; `_supabase_safe` logs Supabase failures.
        """
        if not chatLid or not tipo or not alvo_iso:
            return None
        return self._supabase_safe(
            supabase_repo.schedule_followup,
            chatLid,
            {"tipo": tipo, "horario_iso": alvo_iso, "phone": phone},
        )

    def _mirror_followup_sent(
        self,
        chatLid: str,
        tipo: str,
        job_id: str | None = None,
    ) -> None:
        """Best-effort: mark a scheduled follow-up as sent in Supabase.

        Inputs:
            chatLid: WhatsApp chat identifier.
            tipo: Follow-up type (`1h`, `24h` or `15d`).
            job_id: Optional exact `followup_jobs.id`. When omitted, the
                legacy path marks scheduled rows for the chat and type.

        Side effects:
            PATCHes `followup_jobs` to `status='sent'` and sets `sent_at`.

        Returns:
            None.

        Exceptions/fallbacks:
            Never raises; `_supabase_safe` logs Supabase failures. Falls back
            to chat+tipo matching when `job_id` is unavailable.
        """
        if not chatLid or not tipo:
            return
        if job_id:
            self._supabase_safe(supabase_repo.mark_followup_sent, job_id)
            return
        self._supabase_safe(supabase_repo.mark_followup_sent_by_tipo, chatLid, tipo)

    def _mirror_followup_failed(self, job_id: str, reason: str) -> None:
        """Best-effort: mark one scheduled follow-up job as failed.

        Inputs:
            job_id: Exact `followup_jobs.id`.
            reason: Short operational reason stored in the job payload.

        Side effects:
            PATCHes `followup_jobs.status` from `scheduled` to `failed`.

        Returns:
            None.

        Exceptions/fallbacks:
            Never raises; `_supabase_safe` logs Supabase failures.
        """
        if not job_id:
            return
        self._supabase_safe(supabase_repo.mark_followup_failed, job_id, reason)

    def _mirror_followups_cancelled(self, chatLid: str) -> None:
        """Best-effort: cancel every scheduled follow-up of a chat."""
        if not chatLid:
            return
        self._supabase_safe(supabase_repo.cancel_followup, chat_lid=chatLid)

    # ------------------------------------------------------------------
    # Follow-up automático por inatividade
    # ------------------------------------------------------------------

    def _schedule_followup_timer(
        self,
        phone: str,
        chatLid: str,
        tipo: str,
        delay_seconds: float,
        job_id: str | None = None,
    ) -> bool:
        """Schedule a follow-up timer and optionally bind it to a Supabase job.

        Inputs:
            phone: Destination WhatsApp phone number.
            chatLid: WhatsApp chat identifier.
            tipo: Follow-up type (`1h`, `24h` or `15d`).
            delay_seconds: Seconds until the callback should run.
            job_id: Optional `followup_jobs.id` from Supabase.

        Side effects:
            Creates and starts a `threading.Timer`. Records it in
            `followup_timers` by chat and, when `job_id` is present, in
            `followup_job_timers` so startup cannot arm the same job twice.

        Returns:
            True when a timer was scheduled, False when a live timer for the
            same `job_id` already exists.

        Exceptions/fallbacks:
            Does not raise for duplicate jobs. Timer callback errors are
            handled by `_followup_callback`.
        """
        delay_seconds = max(float(delay_seconds), 0.0)
        timer = None
        with self._timer_lock:
            if job_id:
                existing = self.followup_job_timers.get(job_id)
                if existing and existing.is_alive():
                    print(f"{YELLOW}Follow-up job {job_id} ja tem timer ativo; ignorando duplicata{RESET}")
                    return False
            timer = threading.Timer(
                delay_seconds,
                self._run_followup_timer,
                args=(phone, chatLid, tipo, job_id),
            )
            self.followup_timers.setdefault(chatLid, []).append(timer)
            if job_id:
                self.followup_job_timers[job_id] = timer
        timer.start()
        return True

    def _run_followup_timer(
        self,
        phone: str,
        chatLid: str,
        tipo: str,
        job_id: str | None = None,
    ) -> None:
        """Run a follow-up timer callback after releasing its job-id guard.

        Inputs:
            phone: Destination WhatsApp phone number.
            chatLid: WhatsApp chat identifier.
            tipo: Follow-up type (`1h`, `24h` or `15d`).
            job_id: Optional `followup_jobs.id` associated with this timer.

        Side effects:
            Removes the fired timer from `followup_job_timers`, then delegates
            to `_followup_callback`, which may send a Z-API message and update
            Supabase/local mirror state.

        Returns:
            None.

        Exceptions/fallbacks:
            Does not catch callback failures directly; `_followup_callback`
            handles its own external API errors.
        """
        if job_id:
            with self._timer_lock:
                self.followup_job_timers.pop(job_id, None)
        self._followup_callback(phone, chatLid, tipo, job_id=job_id)

    def _followup_callback(
        self,
        phone: str,
        chatLid: str,
        tipo: str,
        job_id: str | None = None,
    ):
        """Send a follow-up message and persist the resulting job state.

        Inputs:
            phone: Destination WhatsApp phone number.
            chatLid: WhatsApp chat identifier.
            tipo: Follow-up type (`1h`, `24h` or `15d`).
            job_id: Optional `followup_jobs.id`; restored Supabase timers pass
                it so only the exact job is marked sent.

        Side effects:
            May re-arm the timer to the next business window, calls OpenAI to
            generate the message, sends it via Z-API, marks the Supabase job
            as sent, and removes the local JSON mirror entry.

        Returns:
            None.

        Exceptions/fallbacks:
            Swallows OpenAI/Z-API/Supabase errors after logging. If `job_id` is
            absent, falls back to marking scheduled jobs by chat and tipo to
            preserve the legacy runtime path.
        """
        # Garante envio apenas dentro do horário útil (9-12, 13-19, seg-sex)
        agora = datetime.now()
        proximo = _proximo_horario_util(agora)
        if proximo > agora:
            delay = (proximo - agora).total_seconds()
            print(f"{YELLOW}Follow-up {tipo} para {chatLid} fora do horário útil — reagendando para {proximo.strftime('%d/%m %H:%M')} ({int(delay)}s){RESET}")
            self._schedule_followup_timer(phone, chatLid, tipo, delay, job_id=job_id)
            return

        try:
            headers = {"client-token": self.zapi_sec_token}
            llm = ChatOpenAI(model="gpt-4o-mini", api_key=os.environ.get("OPENAI_API_KEY"))

            # Carrega nome e histórico (Supabase = verdade, fallback local)
            lead_info = self._read_lead_info(chatLid)
            nome = (lead_info.get("nome") if lead_info else "") or ""

            instrucao = get_instrucao_followup(tipo, nome)

            historico = self._read_history(chatLid)
            mensagens_recentes = self.extrair_mensagens(historico, chatLid)[-6:]

            resp = llm.invoke([
                SystemMessage(content=get_contexto()),
                *[
                    HumanMessage(content=m["content"]) if m["role"] == "user"
                    else AIMessage(content=m["content"])
                    for m in mensagens_recentes
                ],
                SystemMessage(content=instrucao),
            ])

            url = f"{self.base_url}/send-text"
            requests.post(url, json={"phone": phone, "message": resp.content.strip()}, headers=headers)
            print(f"{GREEN}Follow-up {tipo} enviado para {chatLid}{RESET}")
            self._mirror_followup_sent(chatLid, tipo, job_id=job_id)
        except Exception as e:
            print(f"{RED}Erro no follow-up {tipo} para {chatLid}: {e}{RESET}")
        finally:
            self._limpar_followup_arquivo(chatLid, tipo)

    def agendar_followups(self, phone: str, chatLid: str):
        """Schedule the standard inactivity follow-ups for a chat.

        Inputs:
            phone: Destination WhatsApp phone number.
            chatLid: WhatsApp chat identifier.

        Side effects:
            Reads `necessita_followup` from Supabase with local fallback,
            cancels existing pending timers/jobs, upserts three Supabase
            `followup_jobs` rows (`1h`, `24h`, `15d`), schedules in-memory
            timers using the returned job ids when available, and keeps
            `lead_info.json["followups_agendados"]` as a local mirror.

        Returns:
            None.

        Exceptions/fallbacks:
            Supabase write failures are swallowed by `_supabase_safe`; timers
            still run without `job_id` and use the legacy chat+tipo sent
            marker so current follow-up sending keeps working.
        """
        # Checa se o agente indicou que follow-up é necessário
        try:
            info = self._read_lead_info(chatLid)
            if info.get("necessita_followup") is False:
                print(f"{YELLOW}Follow-up desativado para {chatLid}: {info.get('motivo_followup', '')}{RESET}")
                self.cancelar_followups(chatLid)
                return
        except Exception:
            pass

        self.cancelar_followups(chatLid)

        agora = datetime.now()
        alvo_1h  = agora + timedelta(hours=1)
        alvo_24h = agora + timedelta(hours=24)
        alvo_15d = agora + timedelta(days=15)

        # Persiste horários no lead_info.json (mirror local) e em followup_jobs (Supabase = verdade)
        lead_info_path = os.path.join("chats", chatLid, "lead_info.json")
        info = self._read_json_file(lead_info_path, {})
        info["followups_agendados"] = [
            {"tipo": "1h",  "horario_iso": alvo_1h.isoformat(),  "phone": phone},
            {"tipo": "24h", "horario_iso": alvo_24h.isoformat(), "phone": phone},
            {"tipo": "15d", "horario_iso": alvo_15d.isoformat(), "phone": phone},
        ]
        self._write_json_file(lead_info_path, info)

        job_1h = self._mirror_followup_scheduled(chatLid, "1h",  alvo_1h.isoformat(),  phone)
        job_24h = self._mirror_followup_scheduled(chatLid, "24h", alvo_24h.isoformat(), phone)
        job_15d = self._mirror_followup_scheduled(chatLid, "15d", alvo_15d.isoformat(), phone)

        self._schedule_followup_timer(phone, chatLid, "1h", 3600.0, job_id=(job_1h or {}).get("id"))
        self._schedule_followup_timer(phone, chatLid, "24h", 86400.0, job_id=(job_24h or {}).get("id"))
        self._schedule_followup_timer(phone, chatLid, "15d", 15 * 86400.0, job_id=(job_15d or {}).get("id"))

        print(f"{GREEN}Follow-ups agendados para {chatLid}: 1h às {alvo_1h.strftime('%d/%m %H:%M')} | 24h às {alvo_24h.strftime('%d/%m %H:%M')} | 15d às {alvo_15d.strftime('%d/%m %H:%M')}{RESET}")

    def cancelar_followups(self, chatLid: str):
        """Cancel pending follow-ups for a chat after lead activity.

        Inputs:
            chatLid: WhatsApp chat identifier whose pending follow-ups should
            be cancelled.

        Side effects:
            Cancels in-memory `threading.Timer` objects, removes their
            `followup_job_timers` guards, clears the local
            `lead_info.json["followups_agendados"]` mirror, and marks all
            scheduled Supabase `followup_jobs` for the chat as `cancelled`.

        Returns:
            None.

        Exceptions/fallbacks:
            Supabase failures are swallowed by `_supabase_safe`; local JSON
            cleanup failures are logged by `_limpar_followup_arquivo`.
        """
        timers = self.followup_timers.pop(chatLid, [])
        for t in timers:
            t.cancel()
        if timers:
            timer_ids = {id(t) for t in timers}
            with self._timer_lock:
                for job_id, timer in list(self.followup_job_timers.items()):
                    if id(timer) in timer_ids:
                        self.followup_job_timers.pop(job_id, None)
        if timers:
            print(f"{YELLOW}Follow-ups cancelados para {chatLid} (lead respondeu){RESET}")
        self._limpar_followup_arquivo(chatLid, None)
        self._mirror_followups_cancelled(chatLid)

    def _limpar_followup_arquivo(self, chatLid: str, tipo):
        """Remove follow-up do arquivo. Se tipo=None remove todos."""
        lead_info_path = os.path.join("chats", chatLid, "lead_info.json")
        if not os.path.exists(lead_info_path):
            return
        try:
            info = self._read_json_file(lead_info_path, {})
            agendados = info.get("followups_agendados", [])
            if tipo is None:
                info["followups_agendados"] = []
            else:
                info["followups_agendados"] = [x for x in agendados if x.get("tipo") != tipo]
            self._write_json_file(lead_info_path, info)
        except Exception as e:
            print(f"{RED}Erro ao limpar followup_arquivo para {chatLid}: {e}{RESET}")

    # ------------------------------------------------------------------
    # Início de conversa outbound (lead sem histórico)
    # ------------------------------------------------------------------

    def iniciar_conversa(self, phone: str, chatLid: str, mensagem_personalizada: str = "") -> bool:
        """
        Gera e envia a mensagem de abertura (lead frio) para um lead ainda não abordado.

        Inputs:
            phone: número no formato E.164 sem '+' (ex.: '5511999999999').
            chatLid: identificador do chat ('<phone>@lid').
            mensagem_personalizada: se fornecida, é enviada literalmente (sem chamar o LLM).

        Returns:
            True  -> todas as partes da mensagem foram aceitas pela Z-API e o evento
                     'abertura' foi gravado em history.json + Supabase + follow-ups
                     agendados.
            False -> lead já abordado (guard ja_abordado), OU a Z-API rejeitou pelo
                     menos uma parte da mensagem, OU exceção durante o processo.
                     Em caso de rejeição da Z-API NADA é gravado em histórico
                     (lead permanece elegível para retry).

        Side effects:
            - Em sucesso: append em chats/<chatLid>/history.json, mirror em Supabase
              `messages`, e agendamento de follow-ups de inatividade.
            - Em falha de envio ou exceção: popula `self._last_send_error` com o
              corpo da resposta da Z-API (ou texto da exceção, truncado) para que
              o endpoint /chamar-lead possa surfar a causa ao caller.
        """

        lead_dir = os.path.join("chats", chatLid)

        # Histórico: Supabase = verdade, fallback local.
        historico = self._read_history(chatLid)
        if historico:
            ja_abordado = any(
                item.get("data", {}).get("type") == "abertura"
                for item in historico
            )
            ja_marcado_sem_whatsapp = any(
                item.get("data", {}).get("type") == "numero_sem_whatsapp"
                for item in historico
            )
            tem_mensagem_do_lead = any(
                not item.get("data", {}).get("fromMe", True)
                for item in historico
            )
            if ja_abordado or tem_mensagem_do_lead or ja_marcado_sem_whatsapp:
                print(f"{YELLOW}[iniciar_conversa] {chatLid} já abordado, já tem mensagem do lead, ou já marcado sem WhatsApp — abortando{RESET}")
                return False

        try:
            headers = {"client-token": self.zapi_sec_token}

            # Guard de entregabilidade: número sem conta de WhatsApp nunca
            # recebe a mensagem. A Z-API responde 200/ok otimista mesmo nesse
            # caso, então a checagem é feita ANTES de gerar a abertura via LLM,
            # enviar e gravar histórico — caso contrário o lead seria marcado
            # como chamado sem nunca ter recebido nada (aberturas fantasma).
            if self._numero_tem_whatsapp(phone) is False:
                self._registrar_numero_sem_whatsapp(phone, chatLid)
                self._last_send_error = "numero_sem_whatsapp"
                print(f"{YELLOW}[iniciar_conversa] {chatLid}: número sem WhatsApp — abertura não enviada{RESET}")
                return False

            if mensagem_personalizada:
                mensagem = mensagem_personalizada.strip()
            else:
                llm = ChatOpenAI(model="gpt-4o-mini", api_key=os.environ.get("OPENAI_API_KEY"))

                # Tenta carregar dados do lead já conhecidos (Supabase = verdade)
                contexto_lead = ""
                info = self._read_lead_info(chatLid)
                if info:
                    campos_uteis = {k: v for k, v in info.items() if v and k in ("nome", "empresa", "segmento_mercado")}
                    if campos_uteis:
                        contexto_lead = "Dados conhecidos do lead: " + ", ".join(f"{k}: {v}" for k, v in campos_uteis.items()) + "."
                    descricao_manual = info.get("descricao_manual", "")
                    if descricao_manual:
                        contexto_lead += f" Contexto adicional: {descricao_manual}"

                instrucao = get_instrucao_abertura(contexto_lead)

                resposta = llm.invoke([
                    SystemMessage(content=get_contexto()),
                    SystemMessage(content=instrucao),
                ])

                mensagem = _sanitizar_origem_abertura(
                    resposta.content.strip(),
                    contexto_lead,
                )

            # Envia em múltiplos balões se o LLM quebrou com \n\n.
            # Histórico só é gravado se TODAS as partes forem aceitas pela Z-API
            # (evita poluir history.json com aberturas que nunca foram entregues
            # e bloquear retries via a guard `ja_abordado`).
            partes = [p.strip() for p in mensagem.split("\n\n") if p.strip()] or [mensagem]
            url = f"{self.base_url}/send-text"
            self._last_send_error = None

            for i, parte in enumerate(partes):
                delay = max(1, min(int(len(parte) * 60 / (80 * 5)), 15))
                if i > 0:
                    time.sleep(0.5)
                resp = requests.post(
                    url,
                    json={"phone": phone, "message": parte, "delayTyping": delay},
                    headers=headers,
                    timeout=15,
                )
                ok = False
                body_text = (resp.text or "")[:300]
                if 200 <= resp.status_code < 300:
                    try:
                        body = resp.json()
                        ok = (
                            isinstance(body, dict)
                            and "error" not in body
                            and any(k in body for k in ("messageId", "zaapId", "id"))
                        )
                    except ValueError:
                        ok = False
                if not ok:
                    self._last_send_error = f"HTTP {resp.status_code}: {body_text}"
                    print(f"{RED}[iniciar_conversa] Z-API rejeitou parte {i+1}/{len(partes)} para {chatLid}: {self._last_send_error}{RESET}")
                    return False
                print(f"{GREEN}[iniciar_conversa] Parte {i+1}/{len(partes)} enviada para {chatLid}{RESET}")

            # Grava histórico SOMENTE após todas as partes serem aceitas pela Z-API.
            history_path = os.path.join(lead_dir, "history.json")
            historico = self._read_json_file(history_path, []) if os.path.exists(history_path) else []
            novo_evento = {
                "timestamp": time.time(),
                "data": {
                    "fromMe": True,
                    "phone": phone,
                    "chatLid": chatLid,
                    "text": {"message": mensagem},
                    "type": "abertura",
                }
            }
            historico.append(novo_evento)
            self._write_json_file(history_path, historico)
            self._mirror_message_supabase(chatLid, novo_evento)

            # Agenda follow-ups de inatividade
            self.agendar_followups(phone, chatLid)

            return True

        except Exception as e:
            self._last_send_error = f"exception: {e}"
            print(f"{RED}[iniciar_conversa] Erro ao enviar mensagem de abertura para {chatLid}: {e}{RESET}")
            return False

    def _numero_tem_whatsapp(self, phone: str) -> bool | None:
        """Checa via Z-API `phone-exists` se o número tem conta de WhatsApp.

        Inputs:
            phone: número em dígitos E.164 sem '+'.

        Returns:
            True  -> número tem WhatsApp.
            False -> número confirmadamente NÃO tem WhatsApp.
            None  -> a verificação falhou (rede/HTTP/resposta inesperada).
                     Nesse caso o caller deve seguir com o envio (fail-open),
                     para não bloquear disparos por uma instabilidade da Z-API.
        """
        try:
            resp = requests.get(
                f"{self.base_url}/phone-exists/{phone}",
                headers={"client-token": self.zapi_sec_token},
                timeout=15,
            )
            if resp.status_code != 200:
                return None
            body = resp.json()
        except Exception as e:
            print(f"{YELLOW}[phone-exists] falha ao checar {phone}: {e}{RESET}")
            return None
        if not isinstance(body, dict) or "exists" not in body:
            return None
        return bool(body.get("exists"))

    def _registrar_numero_sem_whatsapp(self, phone: str, chatLid: str) -> None:
        """Grava no histórico (local + Supabase) um marcador de que o número
        não tem WhatsApp.

        É o substituto do evento `abertura` para números sem WhatsApp: deixa
        rastro auditável (`payload.type = 'numero_sem_whatsapp'`) sem marcar o
        lead como abordado. O `messageId` é determinístico, então re-registros
        do mesmo lead fazem upsert na mesma linha em vez de duplicar.
        """
        lead_dir = os.path.join("chats", chatLid)
        os.makedirs(lead_dir, exist_ok=True)
        history_path = os.path.join(lead_dir, "history.json")
        historico = self._read_json_file(history_path, []) if os.path.exists(history_path) else []
        evento = {
            "timestamp": time.time(),
            "data": {
                "fromMe": True,
                "phone": phone,
                "chatLid": chatLid,
                "type": "numero_sem_whatsapp",
                "messageId": f"sysmark:no-wa:{chatLid}",
                "text": {
                    "message": "[sistema] Número sem conta de WhatsApp — "
                               "abertura não enviada (Z-API phone-exists=false)."
                },
            },
        }
        historico.append(evento)
        self._write_json_file(history_path, historico)
        self._mirror_message_supabase(chatLid, evento)

    def migrar_lead_se_necessario(self, phone: str, chatLid: str):
        """
        Quando a Z-API retorna o LID real na primeira resposta do lead,
        verifica se existe uma pasta temporária criada com o telefone
        (ex: 11979611556@lid) e migra tudo para a pasta do LID real.
        """
        import shutil
        import phonenumbers
        from phonenumbers import NumberParseException

        # Candidatos: número completo (com DDI) e número nacional (sem DDI),
        # derivado via libphonenumber para funcionar com qualquer país.
        candidatos = [f"{phone}@lid"]
        try:
            parsed = phonenumbers.parse(f"+{phone}", None)
            if phonenumbers.is_valid_number(parsed):
                nacional = str(parsed.national_number)
                if nacional and nacional != phone:
                    candidatos.append(f"{nacional}@lid")
        except NumberParseException:
            pass

        for lid_temp in candidatos:
            if lid_temp == chatLid:
                continue
            lead_dir_temp = os.path.join("chats", lid_temp)
            if not os.path.exists(lead_dir_temp):
                continue

            print(f"{YELLOW}[migrar_lead] Pasta temporária encontrada: {lid_temp} → {chatLid}{RESET}")

            lead_dir_real = os.path.join("chats", chatLid)
            os.makedirs(lead_dir_real, exist_ok=True)

            # --- Merge lead_info.json ---
            lead_info_temp_path = os.path.join(lead_dir_temp, "lead_info.json")
            lead_info_real_path = os.path.join(lead_dir_real, "lead_info.json")

            info_temp = {}
            info_real = {}

            if os.path.exists(lead_info_temp_path):
                info_temp = self._read_json_file(lead_info_temp_path, {})

            if os.path.exists(lead_info_real_path):
                info_real = self._read_json_file(lead_info_real_path, {})

            followups_temp = info_temp.pop("followups_agendados", [])
            info_merged = {**info_real, **info_temp}
            info_merged["phone"] = phone

            # --- Cancelar timers antigos e reagendar com chatLid real ---
            self.cancelar_followups(lid_temp)

            if followups_temp and info_merged.get("necessita_followup") is not False:
                agora = datetime.now()
                timers = []
                novos_agendados = []
                for item in followups_temp:
                    tipo = item.get("tipo")
                    ph   = item.get("phone", phone)
                    alvo = datetime.fromisoformat(item["horario_iso"])
                    segundos = (alvo - agora).total_seconds()
                    if segundos <= 0:
                        segundos = 5.0
                    t = threading.Timer(segundos, self._followup_callback, args=(ph, chatLid, tipo))
                    timers.append(t)
                    novos_agendados.append({"tipo": tipo, "horario_iso": item["horario_iso"], "phone": ph})
                    t.start()
                if timers:
                    self.followup_timers[chatLid] = timers
                    info_merged["followups_agendados"] = novos_agendados
                    print(f"{GREEN}[migrar_lead] Follow-ups remapeados para {chatLid}{RESET}")

            self._write_json_file(lead_info_real_path, info_merged)

            # --- Merge history.json (histórico antigo na frente) ---
            history_temp_path = os.path.join(lead_dir_temp, "history.json")
            history_real_path = os.path.join(lead_dir_real, "history.json")

            hist_temp = []
            hist_real = []

            if os.path.exists(history_temp_path):
                try:
                    hist_temp = self._read_json_file(history_temp_path, [])
                except Exception:
                    pass

            if os.path.exists(history_real_path):
                try:
                    hist_real = self._read_json_file(history_real_path, [])
                except Exception:
                    pass

            self._write_json_file(history_real_path, hist_temp + hist_real)

            # --- Mirror para Supabase: cancela followups antigos do LID temporário,
            # garante a conversa do LID real, sobe perfil merged e re-upserta as
            # mensagens migradas para o conversation_id correto. Tudo best-effort. ---
            self._mirror_followups_cancelled(lid_temp)
            self._supabase_safe(supabase_repo.upsert_conversation, chatLid, phone=phone)
            self._mirror_lead_info_supabase(chatLid, info_merged)
            for ev in hist_temp:
                self._mirror_message_supabase(chatLid, ev)

            # --- Mapear em memória e deletar pasta temporária ---
            self.phone_to_lid[phone] = chatLid
            shutil.rmtree(lead_dir_temp)
            print(f"{GREEN}[migrar_lead] Migração concluída: {lid_temp} → {chatLid}{RESET}")
            break

    def _restaurar_followups(self):
        """Restore follow-up timers from Supabase during agent startup.

        Inputs:
            None.

        Side effects:
            Reads `followup_jobs` from Supabase via `supabase_repo`, schedules
            timers only for rows still in `status='scheduled'`, records each
            timer by job id to avoid duplicate in-memory timers, updates
            `phone_to_lid`, cancels jobs whose profile has
            `necessita_followup=false`, and marks stale/invalid jobs as
            `failed` in Supabase.

        Returns:
            None.

        Exceptions/fallbacks:
            Supabase failures are swallowed by `_supabase_safe`; startup does
            not fall back to local JSON because Supabase is the source of
            truth for follow-up jobs. Overdue jobs up to 24 hours late are
            still sent after 5 seconds; older jobs are marked `failed`.
        """
        agora_utc = datetime.now(timezone.utc)
        rows = self._supabase_safe(supabase_repo.list_pending_followups)
        if rows is None:
            print(f"{YELLOW}Supabase indisponivel; startup nao restaura follow-ups a partir de JSON local{RESET}")
            return
        if not rows:
            print(f"{GREEN}Nenhum follow-up scheduled pendente no Supabase para restaurar{RESET}")
            return

        if rows:
            grouped: Dict[str, list] = {}
            for r in rows:
                conv = r.get("conversations") or {}
                chat_lid = conv.get("chat_lid")
                if not chat_lid:
                    continue
                grouped.setdefault(chat_lid, []).append(r)

            for chatLid, items in grouped.items():
                profile = self._supabase_safe(supabase_repo.get_lead_profile, chatLid) or {}
                if profile.get("necessita_followup") is False:
                    print(f"{YELLOW}Follow-up ignorado na restauração para {chatLid} (necessita_followup=false){RESET}")
                    self._mirror_followups_cancelled(chatLid)
                    continue
                for item in items:
                    job_id = item.get("id")
                    tipo  = item.get("tipo")
                    phone = item.get("phone") or (item.get("payload") or {}).get("phone")
                    if not phone:
                        if job_id:
                            self._mirror_followup_failed(job_id, "phone ausente no startup")
                        continue
                    target_iso = item.get("target_at")
                    try:
                        alvo = datetime.fromisoformat(target_iso) if target_iso else None
                    except Exception:
                        alvo = None
                    if not alvo:
                        if job_id:
                            self._mirror_followup_failed(job_id, "target_at invalido no startup")
                        continue
                    if alvo.tzinfo is None:
                        alvo = alvo.replace(tzinfo=timezone.utc)
                    else:
                        alvo = alvo.astimezone(timezone.utc)
                    segundos = (alvo - agora_utc).total_seconds()
                    if segundos <= 0:
                        atraso = abs(segundos)
                        if atraso > FOLLOWUP_OVERDUE_GRACE_SECONDS:
                            if job_id:
                                self._mirror_followup_failed(
                                    job_id,
                                    "vencido ha mais de 24h no startup",
                                )
                            print(f"{YELLOW}Follow-up {tipo} para {chatLid} vencido ha {int(atraso)}s; marcado failed{RESET}")
                            continue
                        segundos = 5.0
                        print(f"{YELLOW}Follow-up {tipo} para {chatLid} já passou, disparando em 5s{RESET}")
                    else:
                        print(f"{GREEN}Follow-up {tipo} restaurado para {chatLid}: dispara em {int(segundos)}s{RESET}")
                    self._schedule_followup_timer(phone, chatLid, tipo, segundos, job_id=job_id)

                conv_row = self._supabase_safe(supabase_repo.get_conversation, chatLid) or {}
                ph = conv_row.get("phone")
                if ph:
                    self.phone_to_lid[ph] = chatLid
            return

    # ------------------------------------------------------------------
    # Envio de mensagens
    # ------------------------------------------------------------------

    def send_message(self, phone, chatLid, data) -> dict:

        def __send_message(chatLid, data):
            print(f" {GREEN} > __send_message chamado para {chatLid} {RESET}")
            if chatLid in self.answer_list:
                print(f" {YELLOW} > Respondendo para {chatLid} {RESET}")
                if True:
                    print(f" {BLUE} > Enviando resposta para {chatLid} {RESET}")
                    phone = data.get("phone")
                    url = f"{self.base_url}/send-text"
                    print(f"{GREEN}Url para envio: {url}{RESET}")
                    payload = {
                        "phone": phone,
                        "message": self.get_ai_response(phone, chatLid, data),
                    }
                    headers = {"client-token": self.zapi_sec_token}
                    response = requests.request("POST", url, data=payload, headers=headers)
                else:
                    print(f" {RED}\n\n > Chat {chatLid} não permitido {RESET}")
            else:
                print(f" {RED} > Nenhuma resposta encontrada para {chatLid} {RESET}")

        self.answer_list[chatLid] = data
        text = data.get("text", {})
        texto = text.get("message", "") if isinstance(text, dict) else str(text or "")
        delay = min(3.0 + len(texto) * 0.04, 12.0)
        threading.Timer(delay, __send_message, args=(chatLid, data)).start()

    # ------------------------------------------------------------------
    # Transcrição de áudio
    # ------------------------------------------------------------------

    def _transcrever_audio(self, audio_url: str, mime_type: str = "") -> str:
        """Baixa o áudio e transcreve via OpenAI Whisper."""
        try:
            from openai import OpenAI
            response = requests.get(audio_url, timeout=30)
            response.raise_for_status()

            # Usa o mimeType do payload Z-API (mais confiável que Content-Type HTTP)
            tipo = mime_type or response.headers.get("Content-Type", "")
            suffix = ".ogg"
            if "mp4" in tipo or "m4a" in tipo:
                suffix = ".mp4"
            elif "mpeg" in tipo or "mp3" in tipo:
                suffix = ".mp3"
            elif "wav" in tipo:
                suffix = ".wav"

            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(response.content)
                tmp_path = tmp.name

            try:
                client = OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))
                with open(tmp_path, "rb") as audio_file:
                    transcript = client.audio.transcriptions.create(
                        model="whisper-1",
                        file=audio_file,
                        language="pt",
                    )
                print(f"{GREEN}[transcrever_audio] Transcrição: {transcript.text}{RESET}")
                return transcript.text.strip()
            finally:
                os.unlink(tmp_path)

        except Exception as e:
            print(f"{RED}[transcrever_audio] Erro: {e}{RESET}")
            return ""

    # ------------------------------------------------------------------
    # Análise de imagem
    # ------------------------------------------------------------------

    def _analisar_imagem(self, image_url: str, caption: str = "") -> str:
        """Baixa a imagem e descreve o conteúdo via GPT-4o Vision."""
        try:
            from openai import OpenAI
            client = OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))

            prompt = (
                "Descreva o conteúdo desta imagem de forma objetiva e concisa, "
                "como se estivesse relatando para um assistente de vendas. "
                "Foque em informações relevantes como textos visíveis, produtos, documentos, "
                "capturas de tela ou qualquer dado útil para o contexto de uma conversa comercial."
            )
            if caption:
                prompt += f" A legenda enviada junto com a imagem foi: \"{caption}\"."

            messages_vision = [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {"type": "image_url", "image_url": {"url": image_url}},
                    ],
                }
            ]

            response = client.chat.completions.create(
                model="gpt-4o",
                messages=messages_vision,
                max_tokens=500,
            )
            descricao = response.choices[0].message.content.strip()
            print(f"{GREEN}[analisar_imagem] Descrição: {descricao}{RESET}")
            return descricao

        except Exception as e:
            print(f"{RED}[analisar_imagem] Erro: {e}{RESET}")
            return caption or ""

    # ------------------------------------------------------------------
    # Análise de vídeo
    # ------------------------------------------------------------------

    def _analisar_video(self, video_url: str, caption: str = "") -> str:
        """Baixa o vídeo, extrai frames + áudio via ffmpeg e descreve via GPT-4o Vision + Whisper."""
        tmp_video = None
        tmp_dir = None
        try:
            from openai import OpenAI
            client = OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))

            # Baixa o vídeo
            response = requests.get(video_url, timeout=60)
            response.raise_for_status()

            tipo = response.headers.get("Content-Type", "")
            suffix = ".mp4"
            if "webm" in tipo:
                suffix = ".webm"
            elif "avi" in tipo:
                suffix = ".avi"
            elif "mov" in tipo or "quicktime" in tipo:
                suffix = ".mov"

            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(response.content)
                tmp_video = tmp.name

            tmp_dir = tempfile.mkdtemp()

            # Verifica a duração do vídeo via ffprobe
            LIMITE_FRAMES_SEGUNDOS = 150  # 2min30s
            duracao_result = subprocess.run([
                "ffprobe", "-v", "error",
                "-show_entries", "format=duration",
                "-of", "default=noprint_wrappers=1:nokey=1",
                tmp_video,
            ], capture_output=True, text=True, timeout=30)
            try:
                duracao = float(duracao_result.stdout.strip())
            except (ValueError, TypeError):
                duracao = 0

            # Extrai 1 frame por segundo apenas se o vídeo tiver até 2min30s
            extrair_frames = duracao <= LIMITE_FRAMES_SEGUNDOS
            if not extrair_frames:
                print(f"{YELLOW}[analisar_video] Vídeo com {duracao:.1f}s excede limite de {LIMITE_FRAMES_SEGUNDOS}s, frames não serão extraídos{RESET}")

            if extrair_frames:
                if duracao <= 30:
                    fps = 3
                elif duracao <= 60:
                    fps = 2
                else:
                    fps = 1
                print(f"{YELLOW}[analisar_video] Extraindo frames a {fps}fps (duração: {duracao:.1f}s){RESET}")
                try:
                    subprocess.run([
                        "ffmpeg", "-i", tmp_video,
                        "-vf", f"fps={fps}",
                        "-q:v", "2",
                        os.path.join(tmp_dir, "frame%04d.jpg"),
                        "-y", "-loglevel", "error",
                    ], check=True, timeout=120)
                except Exception as e:
                    print(f"{RED}[analisar_video] Erro ao extrair frames: {e}{RESET}")

            # Extrai o áudio do vídeo
            audio_path = os.path.join(tmp_dir, "audio.mp3")
            try:
                subprocess.run([
                    "ffmpeg", "-i", tmp_video,
                    "-vn", "-ar", "16000", "-ac", "1", "-q:a", "4",
                    audio_path,
                    "-y", "-loglevel", "error",
                ], check=True, timeout=120)
            except Exception as e:
                print(f"{YELLOW}[analisar_video] Erro ao extrair áudio: {e}{RESET}")

            # Coleta os frames extraídos (pode ser vazio se vídeo excedeu o limite)
            MAX_FRAMES = 20
            frames = sorted([
                os.path.join(tmp_dir, f)
                for f in os.listdir(tmp_dir)
                if f.endswith(".jpg")
            ])
            # Se extraiu mais frames do que o limite, faz amostragem uniforme
            if len(frames) > MAX_FRAMES:
                step = len(frames) / MAX_FRAMES
                frames = [frames[int(i * step)] for i in range(MAX_FRAMES)]
                print(f"{YELLOW}[analisar_video] Frames reduzidos para {MAX_FRAMES} por amostragem uniforme{RESET}")

            # Transcreve o áudio via Whisper (se existir)
            transcricao = ""
            if os.path.exists(audio_path) and os.path.getsize(audio_path) > 0:
                try:
                    with open(audio_path, "rb") as audio_file:
                        transcript = client.audio.transcriptions.create(
                            model="whisper-1",
                            file=audio_file,
                            language="pt",
                        )
                    transcricao = transcript.text.strip()
                    print(f"{GREEN}[analisar_video] Transcrição do áudio: {transcricao}{RESET}")
                except Exception as e:
                    print(f"{YELLOW}[analisar_video] Sem áudio ou erro na transcrição: {e}{RESET}")

            # Se não há nem frames nem transcrição, retorna fallback
            if not frames and not transcricao:
                print(f"{RED}[analisar_video] Nenhum frame ou áudio disponível{RESET}")
                return caption or ""

            # Monta o payload multimodal
            content = []
            if frames:
                prompt = (
                    "Analise os frames deste vídeo e descreva o conteúdo de forma objetiva e concisa, "
                    "como se estivesse relatando para um assistente de vendas. "
                    "Foque em informações relevantes como textos visíveis, produtos, demonstrações, "
                    "documentos ou qualquer dado útil para o contexto de uma conversa comercial."
                )
            else:
                prompt = (
                    "O vídeo era longo demais para extrair frames, mas o áudio foi transcrito. "
                    "Com base na transcrição abaixo, descreva o conteúdo de forma objetiva e concisa "
                    "como se estivesse relatando para um assistente de vendas."
                )
            if caption:
                prompt += f" A legenda enviada junto com o vídeo foi: \"{caption}\"."
            if transcricao:
                prompt += f" O áudio do vídeo foi transcrito como: \"{transcricao}\"."
            content.append({"type": "text", "text": prompt})

            for frame_path in frames:
                with open(frame_path, "rb") as f:
                    b64 = base64.b64encode(f.read()).decode("utf-8")
                content.append({
                    "type": "image_url",
                    "image_url": {"url": f"data:image/jpeg;base64,{b64}", "detail": "low"},
                })

            response_ai = client.chat.completions.create(
                model="gpt-4o",
                messages=[{"role": "user", "content": content}],
                max_tokens=700,
            )
            descricao = response_ai.choices[0].message.content.strip()
            print(f"{GREEN}[analisar_video] Descrição: {descricao}{RESET}")
            return descricao

        except subprocess.CalledProcessError as e:
            print(f"{RED}[analisar_video] Erro no ffmpeg: {e}{RESET}")
            return caption or ""
        except Exception as e:
            print(f"{RED}[analisar_video] Erro: {e}{RESET}")
            return caption or ""
        finally:
            if tmp_video and os.path.exists(tmp_video):
                os.unlink(tmp_video)
            if tmp_dir and os.path.exists(tmp_dir):
                import shutil
                shutil.rmtree(tmp_dir, ignore_errors=True)

    def _extrair_documento(self, doc_url: str, mime_type: str = "", caption: str = "") -> str:
        """Baixa e extrai o texto de documentos (DOCX, TXT, etc.)."""
        try:
            response = requests.get(doc_url, timeout=30)
            response.raise_for_status()

            tipo = mime_type or response.headers.get("Content-Type", "")
            suffix = ".docx"
            if "spreadsheet" in tipo or "xlsx" in tipo or "excel" in tipo:
                suffix = ".xlsx"
            elif "csv" in tipo:
                suffix = ".csv"
            elif "plain" in tipo or "txt" in tipo:
                suffix = ".txt"
            elif "pdf" in tipo:
                suffix = ".pdf"

            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(response.content)
                tmp_path = tmp.name

            try:
                conteudo = ""
                if suffix == ".docx":
                    from docx import Document
                    doc = Document(tmp_path)
                    paragrafos = [p.text for p in doc.paragraphs if p.text.strip()]
                    conteudo = "\n".join(paragrafos)
                elif suffix == ".xlsx":
                    import openpyxl
                    wb = openpyxl.load_workbook(tmp_path, data_only=True)
                    partes = []
                    for sheet in wb.worksheets:
                        linhas = []
                        for row in sheet.iter_rows(values_only=True):
                            celulas = [str(c) if c is not None else "" for c in row]
                            if any(c.strip() for c in celulas):
                                linhas.append("\t".join(celulas))
                        if linhas:
                            partes.append(f"[Aba: {sheet.title}]\n" + "\n".join(linhas))
                    conteudo = "\n\n".join(partes)
                elif suffix == ".csv":
                    import csv
                    with open(tmp_path, "r", encoding="utf-8", errors="ignore") as f:
                        reader = csv.reader(f)
                        linhas = ["\t".join(row) for row in reader if any(c.strip() for c in row)]
                    conteudo = "\n".join(linhas)
                elif suffix == ".pdf":
                    import fitz  # pymupdf
                    doc_pdf = fitz.open(tmp_path)
                    partes_pdf = []
                    for i in range(len(doc_pdf)):
                        texto = doc_pdf[i].get_text().strip()
                        if texto:
                            partes_pdf.append(f"[Página {i + 1}]\n{texto}")
                    conteudo = "\n\n".join(partes_pdf)
                    doc_pdf.close()
                elif suffix == ".txt":
                    with open(tmp_path, "r", encoding="utf-8", errors="ignore") as f:
                        conteudo = f.read()

                if conteudo:
                    prefixo = f'[Documento enviado{f": {caption}" if caption else ""}]\n'
                    return prefixo + conteudo.strip()
                return caption or ""
            finally:
                os.unlink(tmp_path)

        except Exception as e:
            print(f"{RED}[_extrair_documento] Erro: {e}{RESET}")
            return caption or ""

    # ------------------------------------------------------------------
    # Extração de mensagens
    # ------------------------------------------------------------------

    def extrair_mensagens(self, data: List[Dict[str, Any]], chatLid: str = "", persist_updates: bool = True) -> List[Dict[str, str]]:
        """
        Extrai as mensagens do histórico no formato esperado pela OpenAI.
        Áudios são transcritos via Whisper e a transcrição é persistida no
        histórico para evitar chamadas repetidas à API.
        """
        mensagens = []

        if not isinstance(data, list):
            print(f"{RED}Erro: data não é uma lista{RESET}")
            return mensagens

        data_ordenada = sorted(data, key=lambda x: x.get("timestamp", 0))
        historico_modificado = False

        # Monta índice messageId → conteúdo para lookup de mensagens citadas
        def _extrair_conteudo_simples(msg_data: dict) -> str:
            text = msg_data.get("text", {})
            if isinstance(text, dict):
                content = text.get("message", "")
            else:
                content = str(text) if text else ""
            if not content:
                audio = msg_data.get("audio", {})
                if isinstance(audio, dict):
                    content = audio.get("transcricao", "")
            if not content:
                image = msg_data.get("image", {})
                if isinstance(image, dict):
                    content = image.get("descricao", "") or image.get("caption", "")
            if not content:
                video = msg_data.get("video", {})
                if isinstance(video, dict):
                    content = video.get("descricao", "") or video.get("caption", "")
            if not content:
                document = msg_data.get("document", {})
                if isinstance(document, dict):
                    content = document.get("conteudo", "") or document.get("caption", "") or ""
            if not content:
                c = msg_data.get("contact")
                if isinstance(c, list):
                    c = c[0] if c else None
                if isinstance(c, dict):
                    content = (c.get("displayName") or "").strip() or (
                        f"contato {c['phones'][0]}" if c.get("phones") else ""
                    )
            return (content or "").strip()

        indice_mensagens: Dict[str, str] = {}
        for _item in data_ordenada:
            _msg_data = _item.get("data", {})
            _msg_id = _msg_data.get("messageId", "")
            if _msg_id:
                _conteudo = _extrair_conteudo_simples(_msg_data)
                if _conteudo:
                    indice_mensagens[_msg_id] = _conteudo

        for item in data_ordenada:
            mensagem_data = item.get("data", {})
            text_data = mensagem_data.get("text", {})

            if isinstance(text_data, dict):
                message_content = text_data.get("message", "")
            else:
                message_content = str(text_data) if text_data else ""

            # Se não há texto, verifica se é mensagem de áudio
            if not message_content:
                audio_data = mensagem_data.get("audio", {})
                if isinstance(audio_data, dict):
                    # Usa transcrição já salva para evitar chamar Whisper novamente
                    message_content = audio_data.get("transcricao", "")
                    if not message_content:
                        audio_url = audio_data.get("audioUrl", "")
                        mime_type = audio_data.get("mimeType", "")
                        if audio_url:
                            print(f"{YELLOW}[extrair_mensagens] Áudio detectado, transcrevendo...{RESET}")
                            message_content = self._transcrever_audio(audio_url, mime_type)
                            if message_content:
                                # Persiste transcrição no item para não retranscrever depois
                                audio_data["transcricao"] = message_content
                                historico_modificado = True

            # Se não há texto nem áudio, verifica se é mensagem de imagem
            if not message_content:
                image_data = mensagem_data.get("image", {})
                if isinstance(image_data, dict):
                    # Usa descrição já salva para evitar chamar a API novamente
                    descricao = image_data.get("descricao", "")
                    if not descricao:
                        image_url = image_data.get("imageUrl", "")
                        caption = image_data.get("caption", "")
                        if image_url:
                            print(f"{YELLOW}[extrair_mensagens] Imagem detectada, analisando...{RESET}")
                            descricao = self._analisar_imagem(image_url, caption)
                            if descricao:
                                image_data["descricao"] = descricao
                                historico_modificado = True
                    if descricao:
                        img_url = image_data.get("imageUrl", "")
                        caption = image_data.get("caption", "")
                        quem = "O lead" if not mensagem_data.get("fromMe", False) else "O assistente"
                        text_part = f"[{quem} enviou uma imagem.{' Legenda: ' + caption + '.' if caption else ''} Análise automática: {descricao}]"
                        if img_url and not mensagem_data.get("fromMe", False):
                            message_content = [
                                {"type": "text", "text": text_part},
                                {"type": "image_url", "image_url": {"url": img_url, "detail": "low"}},
                            ]
                        else:
                            message_content = text_part

            # Se não há texto, áudio nem imagem, verifica se é mensagem de vídeo
            if not message_content:
                video_data = mensagem_data.get("video", {})
                if isinstance(video_data, dict):
                    descricao = video_data.get("descricao", "")
                    if not descricao:
                        video_url = video_data.get("videoUrl", "")
                        caption = video_data.get("caption", "")
                        if video_url:
                            print(f"{YELLOW}[extrair_mensagens] Vídeo detectado, analisando...{RESET}")
                            descricao = self._analisar_video(video_url, caption)
                            if descricao:
                                video_data["descricao"] = descricao
                                historico_modificado = True
                    if descricao:
                        quem = "O lead" if not mensagem_data.get("fromMe", False) else "O assistente"
                        message_content = f"[{quem} enviou um vídeo. Análise automática: {descricao}]"

            # Se não há texto, áudio, imagem nem vídeo, verifica se é documento
            if not message_content:
                doc_data = mensagem_data.get("document", {})
                if isinstance(doc_data, dict):
                    message_content = doc_data.get("conteudo", "")
                    if not message_content:
                        doc_url = doc_data.get("documentUrl", "")
                        mime_type = doc_data.get("mimeType", "")
                        caption = doc_data.get("caption", "")
                        if doc_url:
                            print(f"{YELLOW}[extrair_mensagens] Documento detectado, extraindo...{RESET}")
                            message_content = self._extrair_documento(doc_url, mime_type, caption)
                            if message_content:
                                doc_data["conteudo"] = message_content
                                historico_modificado = True

            # Se não há texto/mídia, verifica se é um card de contato (vCard)
            if not message_content:
                contact = mensagem_data.get("contact")
                contatos = contact if isinstance(contact, list) else ([contact] if isinstance(contact, dict) else [])
                partes = []
                for c in contatos:
                    if not isinstance(c, dict):
                        continue
                    nome = (c.get("displayName") or "").strip()
                    brutos = list(c.get("phones") or []) or re.findall(r"waid=(\d+)", c.get("vCard", "") or "")
                    numeros = [f"+{n}" for n in (normalizar_telefone(str(x)) for x in brutos) if n]
                    if nome and numeros:
                        partes.append(f'Nome exibido no WhatsApp: "{nome}". Telefone(s): {", ".join(numeros)}.')
                    elif numeros:
                        partes.append(f'Telefone(s): {", ".join(numeros)}.')
                    elif nome:
                        partes.append(f'Nome exibido no WhatsApp: "{nome}" (sem número legível).')
                if partes:
                    qtd = "um card de contato" if len(partes) == 1 else f"{len(partes)} cards de contato"
                    message_content = f"[O lead encaminhou {qtd}. {' '.join(partes)}]"

            if not message_content:
                continue

            from_me = mensagem_data.get("fromMe", False)
            role = "assistant" if from_me else "user"

            # Se a mensagem é uma resposta a outra, injeta o contexto da citação
            ref_id = mensagem_data.get("referenceMessageId", "")
            if ref_id and isinstance(message_content, str):
                if ref_id in indice_mensagens:
                    citado = indice_mensagens[ref_id]
                    max_citado = 200
                    if len(citado) > max_citado:
                        citado = citado[:max_citado] + "..."
                    message_content = f'[Em resposta a: "{citado}"]\n{message_content.strip()}'
                else:
                    message_content = f'[Em resposta a uma mensagem anterior]\n{message_content.strip()}'

            if isinstance(message_content, str):
                mensagens.append({"role": role, "content": message_content.strip()})
            else:
                mensagens.append({"role": role, "content": message_content})

        # Salva histórico atualizado se houve nova transcrição
        if historico_modificado and chatLid and persist_updates:
            try:
                file = os.path.join("chats", chatLid, "history.json")
                self._write_json_file(file, data_ordenada)
                print(f"{GREEN}[extrair_mensagens] Transcrições salvas em {file}{RESET}")
            except Exception as e:
                print(f"{RED}[extrair_mensagens] Erro ao salvar transcrições: {e}{RESET}")
            # Mirror para Supabase: re-upsert das mensagens com payload enriquecido.
            # `append_message` é idempotente por (conversation_id, message_id),
            # então sempre atualiza o `payload` quando a mensagem já existe.
            for item in data_ordenada:
                self._mirror_message_supabase(chatLid, item)

        return mensagens

    # ------------------------------------------------------------------
    # Geração de resposta com IA
    # ------------------------------------------------------------------

    def get_ai_response(self, phone, chatLid, data) -> str:

        # Lead respondeu, cancela follow-ups pendentes
        self.cancelar_followups(chatLid)

        self.answer_list[chatLid] = data

        # Interrompe imediatamente se já está processando resposta para este chat
        if chatLid in self.processing_chats:
            self.interrupted_chats.add(chatLid)
            msg_id = data.get("messageId")
            if msg_id:
                self._interrupt_cite_ids[chatLid] = msg_id
            print(f" {YELLOW} > Interrupção imediata sinalizada para {chatLid} {RESET}")

        # --- Helpers de envio / agendamento ---

        def dividir_em_mensagens(texto: str) -> list[str]:
            """
            Divide a resposta em partes separadas por parágrafo (\n\n),
            como balões distintos no WhatsApp.
            """
            partes = [p.strip() for p in texto.split("\n\n") if p.strip()]
            return partes if partes else [texto]

        def calcular_delay_digitacao(texto: str, wpm: int = 80) -> float:
            """
            Calcula delay em segundos simulando digitação.
            Convenção padrão: 1 palavra = 5 caracteres.
            """
            chars = len(texto)
            delay = chars * 60 / (wpm * 5)  # segundos
            return max(1.0, min(delay, 15.0))

        def enviar_mensagem(phone, texto, headers, delay_typing: int = 0, reply_message_id: str = None):
            """Envia uma mensagem de texto pro lead. delay_typing = segundos de 'digitando...' antes do envio (1-15). reply_message_id = ID da mensagem a ser citada."""
            url = f"{self.base_url}/send-text"
            payload = {"phone": phone, "message": texto}
            if delay_typing > 0:
                payload["delayTyping"] = min(delay_typing, 15)
            if reply_message_id:
                payload["messageId"] = reply_message_id
            response = requests.post(url, json=payload, headers=headers)
            print(f"\n\n{GREEN}>Resposta enviada (delayTyping={delay_typing}s): {response.text}{RESET}")

        def obter_msgs_pendentes(historico_raw: list) -> list:
            """
            Retorna as mensagens consecutivas do lead no final do histórico que ainda não foram respondidas.
            Ex: se o lead mandou 3 msgs seguidas sem resposta da IA, retorna essas 3.
            """
            pendentes = []
            for item in reversed(historico_raw):
                msg_data = item.get("data", {})
                if msg_data.get("fromMe"):
                    break  # Encontrou resposta da IA, para
                msg_id = msg_data.get("messageId", "")
                if not msg_id:
                    continue
                texto = ""
                text_field = msg_data.get("text", {})
                if isinstance(text_field, dict):
                    texto = text_field.get("message", "")
                elif text_field:
                    texto = str(text_field)
                if not texto:
                    audio = msg_data.get("audio", {})
                    if isinstance(audio, dict):
                        texto = audio.get("transcricao", "")
                if not texto:
                    continue
                pendentes.append({"id": msg_id, "texto": texto[:150]})
            pendentes.reverse()
            return pendentes

        def decidir_citacoes(partes_resposta: list, historico_raw: list) -> list:
            """
            Quando o lead mandou varias mensagens seguidas, decide qual mensagem
            cada parte da resposta da IA deve citar.
            Retorna lista de messageIds (ou None) alinhada com partes_resposta.
            """
            msgs_pendentes = obter_msgs_pendentes(historico_raw)

            # Se o lead mandou apenas 1 mensagem, sem citacao (fluxo normal)
            if len(msgs_pendentes) <= 1:
                return [None] * len(partes_resposta)

            lista_msgs = "\n".join(
                f'[{m["id"]}]: "{m["texto"]}"'
                for m in msgs_pendentes
            )
            lista_partes = "\n".join(
                f'[PARTE {i+1}]: "{p[:150]}"'
                for i, p in enumerate(partes_resposta)
            )

            prompt = (
                "O lead mandou varias mensagens seguidas no WhatsApp antes da IA responder.\n"
                "A IA vai responder em partes. Voce deve associar cada parte da resposta "
                "a mensagem do lead que ela esta respondendo, para que a IA cite (marque) a mensagem certa.\n\n"
                "Mensagens do lead (em ordem):\n" + lista_msgs + "\n\n"
                "Partes da resposta da IA (em ordem):\n" + lista_partes + "\n\n"
                "REGRAS:\n"
                "- Associe cada PARTE a mensagem do lead que ela responde\n"
                "- Se uma parte nao responde nenhuma mensagem especifica (ex: saudacao, pergunta nova), use NONE\n"
                "- Se so tem 1 parte e 1 mensagem pendente, use NONE (fluxo normal)\n"
                "- A citacao serve para o lead saber qual pergunta dele esta sendo respondida\n\n"
                "Responda EXATAMENTE neste formato (uma linha por parte):\n"
                "PARTE 1: messageId ou NONE\n"
                "PARTE 2: messageId ou NONE\n"
                "..."
            )

            try:
                llm_citacao = ChatOpenAI(model="gpt-4o-mini", api_key=os.environ.get("OPENAI_API_KEY"), temperature=0)
                resp = llm_citacao.invoke([SystemMessage(content=prompt)])
                resultado = resp.content.strip()

                ids_validos = {m["id"] for m in msgs_pendentes}
                citacoes = []
                for linha in resultado.split("\n"):
                    linha = linha.strip()
                    if not linha.startswith("PARTE"):
                        continue
                    # Extrai o valor após ":"
                    valor = linha.split(":", 1)[-1].strip()
                    if valor == "NONE" or valor not in ids_validos:
                        citacoes.append(None)
                    else:
                        citacoes.append(valor)
                        print(f" {BLUE} > Citacao parte {len(citacoes)}: {valor} {RESET}")

                # Preenche com None se faltaram partes
                while len(citacoes) < len(partes_resposta):
                    citacoes.append(None)

                return citacoes[:len(partes_resposta)]
            except Exception as e:
                print(f" {RED} > Erro ao decidir citacoes: {e} {RESET}")
                return [None] * len(partes_resposta)

        def gerar_e_enviar(instrucao: str, mensagens_contexto: list, headers):
            """Chama o LLM com uma instrução adicional e envia o resultado ao lead."""
            llm = ChatOpenAI(model="gpt-4o-mini", api_key=os.environ.get("OPENAI_API_KEY"))
            resposta = llm.invoke([
                SystemMessage(content=get_contexto()),
                *[
                    HumanMessage(content=m["content"]) if m["role"] == "user"
                    else AIMessage(content=m["content"])
                    for m in mensagens_contexto
                ],
                SystemMessage(content=instrucao),
            ])
            partes = dividir_em_mensagens(resposta.content)
            for i, parte in enumerate(partes):
                if chatLid in self.interrupted_chats:
                    print(f" {YELLOW} > {chatLid} enviou nova mensagem, interrompendo resposta {RESET}")
                    return
                if i == 0:
                    delay = int(calcular_delay_digitacao(parte, wpm=400))
                else:
                    time.sleep(0.5)
                    delay = int(calcular_delay_digitacao(parte, wpm=250))
                print(f" {YELLOW} > Delay digitação: {delay}s para parte {i+1}/{len(partes)} {RESET}")
                enviar_mensagem(phone, parte, headers, delay_typing=delay)

        def verificar_consulta_disponibilidade(resultado, headers, mensagens_contexto) -> bool:
            """
            Verifica se o LLM quer consultar horários disponíveis para um dia.
            Busca os slots livres e envia as opções ao lead.
            """
            resposta = resultado["messages"][-1]["content"]
            try:
                limpa = resposta.strip()
                if limpa.startswith("```"):
                    limpa = limpa.split("```")[1]
                    if limpa.startswith("json"):
                        limpa = limpa[4:]
                    limpa = limpa.strip()

                parsed = json.loads(limpa)

                if parsed.get("acao") == "consultar_disponibilidade":
                    data_raw     = parsed.get("data", "")
                    data_reuniao = normalizar_data(data_raw)

                    if not data_reuniao:
                        print(f"{RED}[consultar_disponibilidade] Data inválida: {data_raw}{RESET}")
                        return False

                    alvo = datetime.combine(data_reuniao, datetime.min.time())

                    if alvo.date() <= datetime.now().date():
                        gerar_e_enviar(
                            f"O lead sugeriu {data_reuniao.strftime('%d/%m/%Y')} que já passou ou é hoje. "
                            "Peça de forma natural que escolha uma data futura.",
                            mensagens_contexto, headers
                        )
                        return True

                    if alvo.weekday() >= 5:
                        gerar_e_enviar(
                            f"O lead sugeriu {data_reuniao.strftime('%d/%m/%Y')} que cai em fim de semana. "
                            "Explique de forma natural que as demos acontecem apenas em dias úteis (segunda a sexta) e pergunte qual dia da semana funciona melhor pra ele.",
                            mensagens_contexto, headers
                        )
                        return True

                    _produto = (self._read_lead_info(chatLid) or {}).get("produto_indicado")
                    _calendar_id = get_organizer_for_product(_produto)

                    livres = buscar_horarios_livres(data_reuniao, calendar_id=_calendar_id)
                    data_fmt = data_reuniao.strftime("%d/%m")

                    if livres:
                        opcoes = ", ".join(livres[:3])
                        instrucao = (
                            f"Para o dia {data_fmt} temos disponibilidade às {opcoes}. "
                            "Apresente no máximo 3 opções de forma natural e curta, sem listar todos os horários. Pergunte qual prefere."
                        )
                    else:
                        instrucao = (
                            f"Não há horários disponíveis no dia {data_fmt}. "
                            "Informe o lead de forma natural e sugira que escolha outro dia."
                        )

                    gerar_e_enviar(instrucao, mensagens_contexto, headers)
                    return True

            except (json.JSONDecodeError, ValueError, KeyError):
                pass

            return False

        def verificar_reuniao(resultado, headers, mensagens_contexto) -> bool:
            """
            Verifica se a resposta é um agendamento de reunião/demo.
            Salva em reunioes.json e agenda lembrete 30min antes.
            """
            resposta = resultado["messages"][-1]["content"]
            print(f"{BLUE}[verificar_reuniao] Resposta bruta do LLM: {repr(resposta)}{RESET}")
            try:
                # Limpa markdown se houver
                limpa = resposta.strip()
                if limpa.startswith("```"):
                    limpa = limpa.split("```")[1]
                    if limpa.startswith("json"):
                        limpa = limpa[4:]
                    limpa = limpa.strip()

                print(f"{BLUE}[verificar_reuniao] JSON limpo para parse: {repr(limpa)}{RESET}")
                parsed = json.loads(limpa)
                print(f"{BLUE}[verificar_reuniao] JSON parsed: {parsed}{RESET}")

                if parsed.get("acao") == "agendar_reuniao":
                    data_raw    = parsed.get("data", "")
                    horario_raw = parsed.get("horario", "")
                    mensagem    = parsed.get("mensagem", "Reunião confirmada. Até lá.")
                    email_lead  = parsed.get("email") or None
                    nome_lead   = parsed.get("nome") or None

                    # Fallback: busca email/nome do lead_info (Supabase com fallback local)
                    lead_info = self._read_lead_info(chatLid)
                    email_lead = email_lead or lead_info.get("email") or None
                    nome_lead  = nome_lead  or lead_info.get("nome")  or None

                    # Guard: só agenda se o lead já foi qualificado (produto_indicado definido)
                    produto_indicado = lead_info.get("produto_indicado") or None
                    organizer_email  = get_organizer_for_product(produto_indicado)
                    if not produto_indicado:
                        print(f"{YELLOW}[verificar_reuniao] Lead ainda não qualificado (produto_indicado=null). Bloqueando agendamento.{RESET}")
                        gerar_e_enviar(
                            "O lead tentou agendar uma reunião mas ainda não foi qualificado. Você ainda não identificou se o produto indicado é Squad AI, SaaS Btime ou Ambos. "
                            "Não agende a reunião agora. Continue a conversa fazendo as perguntas de qualificação necessárias para entender o cenário do lead antes de propor a demo.",
                            mensagens_contexto, headers
                        )
                        return True

                    data_reuniao    = normalizar_data(data_raw)
                    horario_reuniao = normalizar_horario(horario_raw)

                    if not data_reuniao or not horario_reuniao:
                        print(f"{RED}Data/horário inválido: {data_raw} {horario_raw}{RESET}")
                        gerar_e_enviar(
                            "Não foi possível interpretar a data ou horário combinado. Peça ao lead, de forma natural, que informe novamente.",
                            mensagens_contexto, headers
                        )
                        return True

                    alvo  = datetime.combine(data_reuniao, datetime.strptime(horario_reuniao, "%H:%M").time())
                    agora = datetime.now()

                    if alvo <= agora:
                        print(f"{RED}Horário de reunião já passou: {alvo}{RESET}")
                        gerar_e_enviar(
                            f"O lead sugeriu uma data/horário que já passou ({alvo.strftime('%d/%m/%Y %H:%M')}). Peça a ele, de forma natural, que escolha uma data futura.",
                            mensagens_contexto, headers
                        )
                        return True

                    # Valida dia útil (seg–sex) e horário comercial (9h–12h ou 13h–19h)
                    if alvo.weekday() >= 5:  # 5=sábado, 6=domingo
                        print(f"{RED}Reunião em fim de semana: {alvo}{RESET}")
                        gerar_e_enviar(
                            f"O lead sugeriu {alvo.strftime('%d/%m/%Y')} que cai em um fim de semana. "
                            "Explique de forma natural que as demos acontecem apenas em dias úteis (segunda a sexta) "
                            "e pergunte qual dia da semana funciona melhor pra ele.",
                            mensagens_contexto, headers
                        )
                        return True

                    hora = alvo.hour + alvo.minute / 60
                    em_manha = 9 <= hora < 12
                    em_tarde = 13 <= hora < 19
                    if not (em_manha or em_tarde):
                        print(f"{RED}Horário fora do comercial: {alvo.strftime('%H:%M')}{RESET}")
                        gerar_e_enviar(
                            f"O lead sugeriu {alvo.strftime('%H:%M')} que está fora do horário disponível (9h às 12h ou 13h às 19h). "
                            "Explique de forma natural e sugira um horário dentro das janelas disponíveis, entre 9h e 12h ou 13h e 19h.",
                            mensagens_contexto, headers
                        )
                        return True

                    # Cancela reunião anterior do mesmo lead, se existir.
                    # Fonte de verdade = Supabase (`list_active_meetings`); o JSON
                    # local é apenas mirror que pode estar incompleto.
                    arquivo_reunioes = "reunioes.json"
                    reunioes = self._read_json_file(arquivo_reunioes, [])

                    reunioes_filtradas = [r for r in reunioes if r.get("chatLid") != chatLid]

                    ativas_no_supabase = self._supabase_safe(
                        supabase_repo.list_active_meetings, chatLid
                    ) or []
                    canceladas_iso = set()
                    for m in ativas_no_supabase:
                        event_id_antigo = m.get("event_id")
                        iso_dt = m.get("scheduled_for")
                        if event_id_antigo:
                            print(f"{YELLOW}Removendo evento anterior (Supabase): {event_id_antigo}{RESET}")
                            deletar_evento_google_calendar(event_id_antigo, calendar_id=organizer_email)
                        if iso_dt:
                            self._mirror_meeting_cancelled(chatLid, iso_dt)
                            canceladas_iso.add(iso_dt)

                    # Fallback: cobre eventuais entradas só locais (Supabase indisponível ou drift)
                    for r in reunioes:
                        if r.get("chatLid") != chatLid:
                            continue
                        iso_dt = r.get("datetime")
                        if iso_dt and iso_dt in canceladas_iso:
                            continue
                        event_id_antigo = r.get("event_id")
                        if event_id_antigo:
                            print(f"{YELLOW}Removendo evento anterior (fallback local): {event_id_antigo}{RESET}")
                            deletar_evento_google_calendar(event_id_antigo, calendar_id=organizer_email)
                        self._mirror_meeting_cancelled(chatLid, iso_dt)

                    if verificar_conflito_google_calendar(data_reuniao, horario_reuniao, calendar_id=organizer_email):
                        livres_no_dia = buscar_horarios_livres(data_reuniao, calendar_id=organizer_email)
                        # Remove o horário conflitante da lista de sugestões
                        livres_no_dia = [h for h in livres_no_dia if h != horario_reuniao]
                        if livres_no_dia:
                            opcoes = ", ".join(livres_no_dia[:3])
                            instrucao_conflito = (
                                f"O horário solicitado ({alvo.strftime('%d/%m/%Y às %H:%M')}) já está ocupado. "
                                f"Informe o lead de forma natural e sugira até 3 opções nesse dia: {opcoes}."
                            )
                        else:
                            instrucao_conflito = (
                                f"O horário solicitado ({alvo.strftime('%d/%m/%Y às %H:%M')}) está ocupado e não há mais horários livres nesse dia. "
                                "Informe o lead de forma natural e peça que escolha outro dia."
                            )
                        gerar_e_enviar(instrucao_conflito, mensagens_contexto, headers)
                        return True

                    meet_link, event_id = criar_evento_google_meet(data_reuniao, horario_reuniao, email_lead=email_lead, nome_lead=nome_lead, rd_email=organizer_email)

                    nova_reuniao = {
                        "phone": phone,
                        "chatLid": chatLid,
                        "datetime": alvo.isoformat(),
                        "agendado_em": agora.isoformat(),
                        "meet_link": meet_link,
                        "event_id": event_id,
                    }
                    reunioes_filtradas.append(nova_reuniao)

                    self._write_json_file(arquivo_reunioes, reunioes_filtradas)
                    self._mirror_meeting_supabase(chatLid, nova_reuniao)

                    data_fmt = alvo.strftime("%d/%m/%Y às %H:%M")
                    link_info = f"Link do Google Meet: {meet_link}" if meet_link else "O convite com o link chegará por e-mail."
                    instrucao_confirmacao = (
                        f"A reunião foi confirmada e criada no Google Calendar. "
                        f"Data: {data_fmt}. {link_info}. "
                        f"Escreva UMA mensagem de confirmação para o lead. "
                        f"Tom: profissional e direto, como uma SDR experiente encerrando o agendamento. "
                        f"Confirme a data/hora, mencione que o convite foi enviado por e-mail e inclua o link se houver. "
                        f"Sem emojis. Sem frases genéricas. Sem JSON, sem markdown, só o texto da mensagem."
                    )
                    gerar_e_enviar(instrucao_confirmacao, mensagens_contexto, headers)

                    # Lembrete 30min antes
                    segundos_total    = (alvo - agora).total_seconds()
                    segundos_lembrete = segundos_total - 1800

                    if segundos_lembrete > 0:
                        link_lembrete = meet_link or ""
                        instrucao_lembrete = get_instrucao_lembrete(horario_reuniao, link_lembrete)
                        def enviar_lembrete_gerado(phone, instrucao, headers):
                            llm_lembrete = ChatOpenAI(model="gpt-4o-mini", api_key=os.environ.get("OPENAI_API_KEY"))
                            resp = llm_lembrete.invoke([
                                SystemMessage(content=get_contexto()),
                                SystemMessage(content=instrucao),
                            ])
                            delay_lembrete = int(calcular_delay_digitacao(resp.content.strip(), wpm=400))
                            enviar_mensagem(phone, resp.content.strip(), headers, delay_typing=delay_lembrete)

                        threading.Timer(segundos_lembrete, enviar_lembrete_gerado, args=(phone, instrucao_lembrete, headers)).start()
                        print(f"{GREEN}Lembrete agendado para {(alvo - timedelta(minutes=30)).strftime('%d/%m %H:%M')}{RESET}")

                    print(f"{GREEN}Reunião agendada: {alvo.strftime('%d/%m/%Y %H:%M')} | Meet: {meet_link}{RESET}")
                    return True

            except (json.JSONDecodeError, ValueError, KeyError) as e:
                print(f"{YELLOW}[verificar_reuniao] Resposta não é JSON de reunião: {type(e).__name__}: {e}{RESET}")

            return False

        def verificar_agendamento(resultado, headers) -> bool:
            """
            Verifica se a resposta é um agendamento de retorno.
            Retorna True se agendou (para parar o fluxo normal), False caso contrário.
            """
            resposta = resultado["messages"][-1]["content"]
            try:
                limpa = resposta.strip()
                if limpa.startswith("```"):
                    limpa = limpa.split("```")[1]
                    if limpa.startswith("json"):
                        limpa = limpa[4:]
                    limpa = limpa.strip()

                parsed = json.loads(limpa)

                if parsed.get("acao") == "agendar_retorno":
                    horario_raw = parsed.get("horario", "")

                    horario = normalizar_horario(horario_raw)
                    if not horario:
                        print(f"{RED}Horário não reconhecido: {horario_raw}{RESET}")
                        return False

                    agora = datetime.now()
                    alvo = datetime.strptime(horario, "%H:%M").replace(
                        year=agora.year, month=agora.month, day=agora.day
                    )

                    if alvo <= agora:
                        alvo += timedelta(days=1)
                        print(f"{YELLOW}Horário já passou, agendando para amanhã: {alvo}{RESET}")

                    segundos = (alvo - agora).total_seconds()

                    gerar_e_enviar(
                        f"O lead pediu para ser contatado às {horario}. "
                        f"Confirme de forma natural que você retornará nesse horário. "
                        f"Sem emojis. Sem JSON. Só o texto da mensagem.",
                        [],
                        headers,
                    )
                    threading.Timer(segundos, create_answer, args=(phone, chatLid, data), kwargs={"is_recontato": True}).start()
                    print(f"{GREEN}Recontato agendado para {alvo.strftime('%d/%m %H:%M')} ({int(segundos)}s){RESET}")

                    # Retorno combinado → desativa follow-ups de inatividade
                    self.cancelar_followups(chatLid)
                    lead_info_path = os.path.join("chats", chatLid, "lead_info.json")
                    try:
                        info_ret = self._read_lead_info(chatLid)
                        info_ret["necessita_followup"] = False
                        info_ret["motivo_followup"] = f"Retorno combinado agendado para {alvo.strftime('%d/%m %H:%M')}, follow-up de inatividade não necessário"
                        self._write_json_file(lead_info_path, info_ret)
                        self._mirror_lead_info_supabase(chatLid, info_ret)
                    except Exception as e:
                        print(f"{RED}Erro ao desativar follow-up para retorno combinado: {e}{RESET}")

                    return True

            except (json.JSONDecodeError, ValueError, KeyError):
                pass

            return False

        def atualizar_info_lead(chatLid, mensagens, ultima_resposta):
            """
            Usa o LLM para extrair e atualizar informações estruturadas do lead
            com base na conversa até o momento. Salva em chats/{chatLid}/lead_info.json.
            """
            try:
                llm = ChatOpenAI(model="gpt-4o-mini", api_key=os.environ.get("OPENAI_API_KEY"))

                prompt_extracao = get_prompt_extracao_lead()

                conversa_texto = "\n".join(
                    f"{'Lead' if m['role'] == 'user' else 'SDR'}: {m['content']}"
                    for m in mensagens
                ) + f"\nSDR: {ultima_resposta}"

                resposta = llm.invoke([
                    SystemMessage(content=prompt_extracao),
                    HumanMessage(content=conversa_texto),
                ])

                limpa = resposta.content.strip()
                if limpa.startswith("```"):
                    limpa = limpa.split("```")[1]
                    if limpa.startswith("json"):
                        limpa = limpa[4:]
                    limpa = limpa.strip()

                info  = json.loads(limpa)
                info["atualizado_em"] = datetime.now().isoformat()

                lead_dir    = os.path.join("chats", chatLid)
                os.makedirs(lead_dir, exist_ok=True)
                arquivo_info = os.path.join(lead_dir, "lead_info.json")

                if os.path.exists(arquivo_info):
                    info_anterior = self._read_json_file(arquivo_info, {})
                    # Campos que sempre devem refletir o estado mais recente da conversa
                    _sempre_atualizar = {"necessita_followup", "motivo_followup", "estagio_conversa", "atualizado_em"}
                    _nunca_sobrescrever = {"tentativas_retomada"}
                    for key, val in info_anterior.items():
                        if key in _sempre_atualizar:
                            continue
                        if key in _nunca_sobrescrever:
                            info[key] = val  # sempre preserva o valor gerenciado pelo sistema
                            continue
                        if info.get(key) in (None, [], "") and val not in (None, [], ""):
                            # novo valor é vazio → mantém o anterior para não perder info
                            info[key] = val
                        # se o novo valor não é vazio, usa o novo (permite corrigir dados errados)

                self._write_json_file(arquivo_info, info)
                self._mirror_lead_info_supabase(chatLid, info)

                print(f"{GREEN}Lead info atualizado: {arquivo_info}{RESET}")

            except Exception as e:
                print(f"{RED}Erro ao atualizar lead_info: {e}{RESET}")

        # --- Orquestração principal ---

        def create_answer(phone, chatLid, data, history_limit=50, is_recontato=False):
            self.pending_timers.pop(chatLid, None)
            self.pending_texts.pop(chatLid, None)

            # Pega messageId da mensagem que causou interrupção (setado em get_ai_response)
            interrupt_reply_id = self._interrupt_cite_ids.pop(chatLid, None)

            # Se já está processando resposta, interromper a anterior
            if chatLid in self.processing_chats:
                self.interrupted_chats.add(chatLid)
                if not interrupt_reply_id:
                    interrupt_reply_id = data.get("messageId")
                print(f" {YELLOW} > Interrompendo resposta anterior para {chatLid} {RESET}")
                # Espera a resposta anterior parar antes de continuar
                while chatLid in self.processing_chats:
                    time.sleep(0.2)

            # Limpa estado de interrupção antes de iniciar novo processamento
            self.interrupted_chats.discard(chatLid)

            print(f" {GREEN} > get_ai_response chamado para {chatLid} {RESET}")
            print(f"{RED}Anser_list: {self.answer_list}{RESET}")

            self.processing_chats.add(chatLid)
            try:
                _create_answer_body(phone, chatLid, data, history_limit, is_recontato, interrupt_reply_id=interrupt_reply_id)
            except Exception as e:
                print(f"{RED}Erro inesperado em create_answer para {chatLid}: {e}{RESET}")
            finally:
                self.processing_chats.discard(chatLid)

        def _create_answer_body(phone, chatLid, data, history_limit=50, is_recontato=False, interrupt_reply_id=None):

            if is_recontato:
                headers = {"client-token": self.zapi_sec_token}
                llm = ChatOpenAI(model="gpt-4o-mini", api_key=os.environ.get("OPENAI_API_KEY"))

                historico_recontato = self._read_history(chatLid)
                mensagens_recontato = self.extrair_mensagens(historico_recontato, chatLid)[-10:]

                resposta_recontato = llm.invoke([
                    SystemMessage(content=get_contexto()),
                    *[
                        HumanMessage(content=m["content"]) if m["role"] == "user"
                        else AIMessage(content=m["content"])
                        for m in mensagens_recontato
                    ],
                    SystemMessage(content=get_instrucao_recontato()),
                ])
                partes_recontato = dividir_em_mensagens(resposta_recontato.content)
                for i, parte in enumerate(partes_recontato):
                    if chatLid in self.interrupted_chats:
                        print(f" {YELLOW} > {chatLid} enviou nova mensagem, interrompendo recontato {RESET}")
                        return
                    if i == 0:
                        delay = int(calcular_delay_digitacao(parte, wpm=400))
                    else:
                        time.sleep(0.5)
                        delay = int(calcular_delay_digitacao(parte, wpm=250))
                    print(f" {YELLOW} > Delay digitação (recontato): {delay}s para parte {i+1}/{len(partes_recontato)} {RESET}")
                    enviar_mensagem(phone, parte, headers, delay_typing=delay)
                self.answer_list[phone] = data
                return

            if chatLid in self.answer_list:
                print(f" {YELLOW} > Respondendo para {chatLid} {RESET}")

                if True:
                    print(f" {BLUE} > Processo de resposta iniciado.\nChatLid: {chatLid} {RESET}")

                    # Histórico: Supabase = verdade; fallback para arquivo local em caso de falha.
                    historico = self._read_history(chatLid) or None
                    if historico is None:
                        file = os.path.join("chats", chatLid, "history.json")
                        if os.path.exists(file):
                            for _tentativa in range(3):
                                try:
                                    with open(file, "r", encoding="utf-8") as f:
                                        historico = json.load(f)
                                    break
                                except (json.JSONDecodeError, ValueError, OSError):
                                    # Arquivo pode estar sendo escrito concorrentemente; aguarda e tenta de novo
                                    time.sleep(0.3)
                    if historico is None:
                        historico = [{"data": data, "timestamp": data.get("momment", 0)}]

                    # Separa documentos (sempre incluídos) do restante (limitado por history_limit)
                    docs = [item for item in historico if isinstance(item.get("data", {}).get("document"), dict)]
                    nao_docs = [item for item in historico if not isinstance(item.get("data", {}).get("document"), dict)]
                    historico_filtrado = sorted(
                        docs + nao_docs[-history_limit:],
                        key=lambda x: x.get("timestamp", 0)
                    )
                    mensagens_formatadas = self.extrair_mensagens(historico_filtrado, chatLid, persist_updates=False)

                    # Injeta dados já conhecidos do lead no contexto (Supabase = verdade)
                    contexto = get_contexto()
                    lead_info_path = os.path.join("chats", chatLid, "lead_info.json")
                    lead_info_atual = self._read_lead_info(chatLid)
                    if lead_info_atual:
                        campos = {
                            "nome":  lead_info_atual.get("nome"),
                            "email": lead_info_atual.get("email"),
                        }
                        conhecidos = {k: v for k, v in campos.items() if v}
                        if conhecidos:
                            linhas = ", ".join(f"{k}: {v}" for k, v in conhecidos.items())
                            contexto += f"\n\nDados já conhecidos deste lead (não peça novamente): {linhas}."

                    # Injeta aviso se este lead foi abordado de forma outbound (lead frio)
                    tem_abertura = any(
                        item.get("data", {}).get("type") == "abertura"
                        for item in historico
                    )
                    if tem_abertura:
                        contexto += (
                            "\n\nATENÇÃO: você já se apresentou para este lead (abordagem outbound). "
                            "NÃO se apresente novamente. Retome a conversa de onde parou."
                        )

                    # Injeta aviso se já existe reunião agendada para este lead (Supabase = verdade)
                    try:
                        reuniao_lead = self._read_active_meeting_for(chatLid)
                        if reuniao_lead:
                            dt_reuniao = reuniao_lead.get("datetime", "")
                            try:
                                alvo_reuniao = datetime.fromisoformat(dt_reuniao) if dt_reuniao else None
                                # Normaliza para comparação naive (datetime.now()) caso venha com tz do Supabase
                                if alvo_reuniao and alvo_reuniao.tzinfo is not None:
                                    alvo_reuniao = alvo_reuniao.astimezone().replace(tzinfo=None)
                            except Exception:
                                alvo_reuniao = None
                            if alvo_reuniao and alvo_reuniao > datetime.now():
                                dt_fmt = alvo_reuniao.strftime("%d/%m/%Y às %H:%M")
                                meet_link_exist = reuniao_lead.get("meet_link", "")
                                contexto += (
                                    f"\n\nATENÇÃO — REUNIÃO JÁ AGENDADA: A demo com este lead está marcada para {dt_fmt}."
                                    f"{' Link: ' + meet_link_exist if meet_link_exist else ''} "
                                    f"NÃO gere um novo JSON de agendamento. "
                                    f"REGRAS OBRIGATÓRIAS para este estado pós-agendamento:\n"
                                    f"- Se o lead enviar mensagem de encerramento (ex: 'Obrigada', 'Ok', 'Até lá'), responda com despedida curta e profissional.\n"
                                    f"- Se o lead enviar emojis, figurinhas ou mensagens sem texto claro, responda de forma breve e natural, como em uma conversa casual. NUNCA pergunte se quer remarcar nesses casos.\n"
                                    f"- Se o lead enviar mensagem confusa ou fora de contexto, peça esclarecimento de forma natural e direta. NÃO assuma que ele quer remarcar.\n"
                                    f"- Só mencione reagendamento ou cancelamento se o lead EXPLICITAMENTE pedir isso.\n"
                                    f"- NUNCA inicie resposta com 'Claro.' seguido de pergunta sobre remarcar. Isso soa robótico e confunde o lead.\n"
                                    f"- Não repita as informações da reunião a menos que o lead pergunte."
                                )
                                print(f"{GREEN}[create_answer] Reunião já agendada para {chatLid} em {dt_fmt}, injetando aviso no contexto.{RESET}")
                    except Exception as e:
                        print(f"{RED}[create_answer] Erro ao verificar reunião existente: {e}{RESET}")

                    # Injeta horários livres dos próximos 3 dias úteis no contexto
                    try:
                        from datetime import date
                        _produto_ctx = lead_info_atual.get("produto_indicado") if lead_info_atual else None
                        _calendar_ctx = get_organizer_for_product(_produto_ctx)
                        horarios_por_dia = {}
                        dia = datetime.now().date() + timedelta(days=1)
                        max_busca_dias = 21
                        dias_verificados = 0
                        while len(horarios_por_dia) < 3 and dias_verificados < max_busca_dias:
                            if dia.weekday() < 5:  # seg–sex
                                livres = buscar_horarios_livres(dia, calendar_id=_calendar_ctx)
                                if livres:
                                    horarios_por_dia[dia.strftime("%d/%m/%Y (%A)")] = livres
                            dia += timedelta(days=1)
                            dias_verificados += 1
                        if horarios_por_dia:
                            linhas_agenda = "\n".join(
                                f"- {d}: {', '.join(slots[:3])}"
                                for d, slots in horarios_por_dia.items()
                            )
                            contexto += f"\n\nHorários disponíveis na agenda (sugira apenas 2 ou 3 opções de forma natural, não liste tudo):\n{linhas_agenda}"
                    except Exception as e:
                        print(f"{RED}Erro ao buscar horários livres para contexto: {e}{RESET}")

                    # --- Micro agentes (paralelo) + AnaAgent ---
                    # Encontra a última mensagem do user (pode não ser a última do histórico
                    # se a AI enviou msg pós-interrupção)
                    lead_message_atual = ""
                    last_user_idx = -1
                    for _i in range(len(mensagens_formatadas) - 1, -1, -1):
                        if mensagens_formatadas[_i].get("role") == "user" and mensagens_formatadas[_i].get("content"):
                            lead_message_atual = mensagens_formatadas[_i]["content"]
                            last_user_idx = _i
                            break

                    if not lead_message_atual:
                        print(f"{RED}[create_answer] Nenhuma mensagem do lead encontrada para {chatLid}, abortando{RESET}")
                        return

                    # History = tudo antes da última msg do user
                    # (exclui msgs AI pós-interrupção que não devem influenciar a nova resposta)
                    history_for_agent = mensagens_formatadas[:last_user_idx] if last_user_idx > 0 else []

                    # Detecção estrutural de vCard: última mensagem do lead é só contato (sem texto/áudio/imagem)?
                    lead_message_eh_vcard = False
                    for _item in reversed(historico_filtrado):
                        _md = _item.get("data", {}) or {}
                        if _md.get("fromMe", False):
                            continue
                        _text = (_md.get("text") or {}).get("message") or ""
                        _audio_transc = (_md.get("audio") or {}).get("transcricao") or ""
                        lead_message_eh_vcard = bool(_md.get("contact")) and not _text and not _audio_transc and not _md.get("image")
                        break

                    print(f"{BLUE}[micro_agents] Iniciando análise paralela para {chatLid}{RESET}")
                    micro_context = run_micro_agents(
                        lead_message_atual,
                        history_for_agent,
                        has_active_meeting=self._has_active_meeting(chatLid),
                        lead_message_eh_vcard=lead_message_eh_vcard,
                    )
                    print(f"{BLUE}[micro_agents] Concluído: {micro_context}{RESET}")

                    _closure = micro_context.get("closure", {})
                    if _closure.get("should_respond") is False and _closure.get("confidence") == "high":
                        print(f"{YELLOW}[closure] Silenciando {chatLid} — {_closure.get('reason')}{RESET}")
                        self.pending_texts.pop(chatLid, None)
                        self.answer_list.pop(chatLid, None)
                        return

                    # --- Handoff: lead repassou contato de outra pessoa ---
                    _handoff_instrucao_pos_disparo = ""
                    _handoff = micro_context.get("handoff", {})
                    if (
                        _handoff.get("is_handoff")
                        and _handoff.get("confidence") == "high"
                        and _handoff.get("new_phone")
                    ):
                        novo_numero = _handoff["new_phone"]
                        novo_nome = _handoff.get("new_contact_name") or "Responsável"
                        _li = self._read_lead_info(chatLid)
                        phone_atual = _li.get("phone", "")

                        if novo_numero and novo_numero == phone_atual:
                            # Lead passou o próprio número como se fosse de outra pessoa.
                            # Não dispara handoff, mas avisa a Ana pra pedir esclarecimento
                            # em vez de alucinar que vai chamar.
                            _handoff_instrucao_pos_disparo = (
                                f"\n\nATENÇÃO: o lead indicou o número {novo_numero} como contato "
                                f"de {_handoff.get('new_contact_name') or 'outra pessoa'}, mas esse "
                                f"é o MESMO número com o qual você está conversando agora. "
                                f"NÃO confirme que vai chamar ninguém. Aponte essa coincidência de "
                                f"forma natural e peça o número correto do outro contato. "
                                f"NÃO emita JSON ou qualquer formato estruturado — apenas texto curto."
                            )
                        elif novo_numero and novo_numero != phone_atual:
                            empresa_orig = _li.get("empresa", "")
                            segmento_orig = _li.get("segmento_mercado", "")
                            produto_orig = _li.get("produto_indicado", "")
                            desafios_orig = _li.get("desafios_identificados") or []
                            indicado_por = _li.get("nome", "")
                            qual = micro_context.get("qualification", {}) or {}

                            resumo_partes = []
                            if empresa_orig:
                                resumo_partes.append(f"empresa {empresa_orig}")
                            if segmento_orig:
                                resumo_partes.append(f"segmento {segmento_orig}")
                            if qual.get("revenue_value"):
                                resumo_partes.append(f"faturamento ~{qual['revenue_value']}")
                            pain = qual.get("pain_description") or (
                                "; ".join(desafios_orig) if desafios_orig else ""
                            )
                            if pain:
                                resumo_partes.append(f"dor: {pain}")
                            produto = qual.get("product_routed") or produto_orig
                            if produto:
                                resumo_partes.append(f"produto: {produto}")
                            resumo = "; ".join(resumo_partes) or "sem contexto adicional"

                            descricao = f"Indicado por {indicado_por or 'contato anterior'}"
                            if empresa_orig and indicado_por:
                                descricao = f"Indicado por {indicado_por} (da {empresa_orig})"
                            if _handoff.get("new_contact_role"):
                                descricao += f" como {_handoff['new_contact_role']}"
                            descricao += f". {resumo}."

                            # Fonte da verdade: tabela `handoffs` no Supabase (success=true).
                            ja_confirmado = bool(self._supabase_safe(
                                supabase_repo.has_successful_handoff,
                                chatLid,
                                novo_numero,
                            ))
                            with self._handoff_in_flight_lock:
                                in_flight_chat = self._handoff_in_flight.setdefault(chatLid, set())
                                em_voo = novo_numero in in_flight_chat

                            if ja_confirmado:
                                print(f"{YELLOW}[handoff] {novo_numero} já confirmado anteriormente — não redispara{RESET}")
                                _handoff_instrucao_pos_disparo = (
                                    f"\n\nO contato de {novo_nome} ({novo_numero}) JÁ FOI acionado "
                                    f"com sucesso anteriormente — não dispare de novo nem prometa "
                                    f"chamar agora. Nesta resposta: explique de forma natural e curta "
                                    f"que o time já foi avisado e está aguardando o {novo_nome} "
                                    f"responder, e mantenha a porta aberta. NÃO emita JSON ou "
                                    f"qualquer formato estruturado — apenas texto natural curto."
                                )
                            elif em_voo:
                                print(f"{YELLOW}[handoff] {novo_numero} já está em voo neste chat — aguardando confirmação{RESET}")
                                _handoff_instrucao_pos_disparo = (
                                    f"\n\nO contato de {novo_nome} ({novo_numero}) está sendo acionado "
                                    f"neste momento (disparo em andamento). Nesta resposta: peça um instante "
                                    f"de forma natural e curta. NÃO emita JSON ou qualquer formato "
                                    f"estruturado — apenas texto natural curto."
                                )
                            else:
                                # Reserva slot in-flight + grava attempt na DB (success=false).
                                # Source of truth = tabela handoffs. Lista local fica como fallback.
                                with self._handoff_in_flight_lock:
                                    self._handoff_in_flight[chatLid].add(novo_numero)

                                _handoff_id = self._supabase_safe(
                                    supabase_repo.record_handoff_attempt,
                                    chatLid,
                                    novo_numero,
                                    novo_nome,
                                    _handoff.get("new_contact_role"),
                                    descricao,
                                )

                                # Slots estruturados pra abertura — sem reparsear free-text.
                                # Pain prioriza qualification.pain_description; fallback pra
                                # desafios_identificados; produto idem.
                                _pain_summary = qual.get("pain_description") or (
                                    "; ".join(desafios_orig) if desafios_orig else None
                                )
                                _produto = qual.get("product_routed") or produto_orig or None
                                _abertura_slots = {
                                    "new_contact_name": novo_nome,
                                    "new_contact_role": _handoff.get("new_contact_role"),
                                    "referrer_name": indicado_por or None,
                                    "referrer_company": empresa_orig or None,
                                    "pain_summary": _pain_summary,
                                    "product": _produto,
                                }
                                _abertura = gerar_abertura_handoff(_abertura_slots)

                                def _disparar_handoff_async(
                                    _novo_nome=novo_nome,
                                    _novo_numero=novo_numero,
                                    _descricao=descricao,
                                    _chatLid=chatLid,
                                    _handoff_id=_handoff_id,
                                    _abertura=_abertura,
                                ):
                                    """Dispara chamar_lead_interno em thread separada para não bloquear a Ana.

                                    Inputs: closures sobre dados do handoff + id da linha em handoffs.
                                    Side effects:
                                      - Chama db_app.chamar_lead_interno (Z-API).
                                      - SUCESSO: marca handoffs.success=true (DB).
                                      - FALHA: marca handoffs.error_detail mantendo success=false
                                        e libera _handoff_in_flight pra próximo turno retry.
                                    """
                                    sucesso = False
                                    erro_detalhe = None
                                    try:
                                        from db_app import chamar_lead_interno
                                        _body, _status = chamar_lead_interno(
                                            nome=_novo_nome,
                                            numero=_novo_numero,
                                            descricao=_descricao,
                                            mensagem=_abertura,
                                            indicado_por_chatLid=_chatLid,
                                        )
                                        if _status == 200 and isinstance(_body, dict) and _body.get("status") == "ok":
                                            sucesso = True
                                            print(f"{GREEN}[handoff] Disparado para {_novo_numero} ({_novo_nome}){RESET}")
                                        else:
                                            erro_detalhe = f"status={_status} body={_body}"
                                            print(f"{RED}[handoff] Disparo não confirmado para {_novo_numero}: {erro_detalhe}{RESET}")
                                    except Exception as e:
                                        erro_detalhe = str(e)
                                        print(f"{RED}[handoff] Falhou disparo para {_novo_numero}: {e}{RESET}")

                                    try:
                                        if sucesso:
                                            if _handoff_id:
                                                self._supabase_safe(supabase_repo.mark_handoff_success, _handoff_id)
                                        else:
                                            if _handoff_id:
                                                self._supabase_safe(
                                                    supabase_repo.mark_handoff_failure,
                                                    _handoff_id,
                                                    erro_detalhe,
                                                )
                                    finally:
                                        with self._handoff_in_flight_lock:
                                            self._handoff_in_flight.get(_chatLid, set()).discard(_novo_numero)
                                        if not sucesso:
                                            print(f"{YELLOW}[handoff] {_novo_numero} liberado para retry no próximo turno{RESET}")

                                threading.Thread(target=_disparar_handoff_async, daemon=True).start()

                                _handoff_instrucao_pos_disparo = (
                                    f"\n\nO contato de {novo_nome} foi recebido e já está sendo "
                                    f"acionado pelo sistema em paralelo. Nesta resposta para a pessoa "
                                    f"atual: apenas agradeça brevemente o repasse de forma natural e "
                                    f"mantenha a porta aberta. NÃO continue qualificando, NÃO peça novas "
                                    f"informações sobre {novo_nome}, e NÃO emita JSON ou qualquer formato "
                                    f"estruturado nesta resposta — apenas texto natural curto."
                                )

                    # --- Tentativas de retomada de lead desinteressado ---
                    _intent_atual = micro_context.get("intent", {}).get("intent", "")
                    _sinais_positivos = {"buy_signal", "interest", "scheduling_request"}
                    _li_path = os.path.join("chats", chatLid, "lead_info.json")
                    _li_data = self._read_lead_info(chatLid)

                    if _intent_atual in _sinais_positivos:
                        if _li_data.get("tentativas_retomada", 0) > 0:
                            print(f"{GREEN}[retomada] Lead re-engajou. Resetando tentativas_retomada para 0.{RESET}")
                        _li_data["tentativas_retomada"] = 0
                    elif _intent_atual == "objection":
                        _li_data["tentativas_retomada"] = _li_data.get("tentativas_retomada", 0) + 1
                        print(f"{YELLOW}[retomada] Lead desinteressado. Tentativa {_li_data['tentativas_retomada']}/4.{RESET}")

                    os.makedirs(os.path.join("chats", chatLid), exist_ok=True)
                    self._write_json_file(_li_path, _li_data)
                    self._mirror_lead_info_supabase(chatLid, _li_data)

                    _tentativas_retomada = _li_data.get("tentativas_retomada", 0)
                    if _tentativas_retomada > 0:
                        if _tentativas_retomada >= 4:
                            contexto += (
                                f"\n\nEste lead demonstrou desinteresse e você já tentou retomá-lo "
                                f"{_tentativas_retomada} vezes. Esta é a última tentativa permitida. "
                                f"Se não houver abertura genuína, encerre a conversa com elegância e respeito."
                            )
                        else:
                            contexto += (
                                f"\n\nEste lead demonstrou desinteresse. "
                                f"Tentativas de retomada realizadas: {_tentativas_retomada} de 4. "
                                f"Continue tentando reverter o desinteresse antes de desistir."
                            )
                    # ----------------------------------------------------

                    # Primeira mensagem do lead: instrução obrigatória de apresentação
                    if not history_for_agent:
                        contexto += (
                            "\n\nINSTRUÇÃO OBRIGATÓRIA — PRIMEIRA MENSAGEM INBOUND: O lead entrou em contato por conta própria. "
                            "Você DEVE obrigatoriamente se apresentar nesta resposta. "
                            "Diga seu nome (Ana) e que é da Btime. "
                            "Como o objetivo do lead é desconhecido, NUNCA assuma o motivo do contato. "
                            "Após a apresentação, pergunte de forma natural como pode ajudá-lo (ex: 'como posso te ajudar?', 'com o que posso te ajudar?'). "
                            "Nunca pule a apresentação. Nunca faça perguntas sobre dores ou segmento antes de entender o que o lead quer."
                        )
                        print(f"{BLUE}[create_answer] Primeira mensagem de {chatLid}: instrução de apresentação injetada.{RESET}")

                    history_lc = _to_langchain_messages(history_for_agent)
                    ana_response = run_ana_agent(
                        lead_message=lead_message_atual,
                        history=history_lc,
                        micro_agent_context=micro_context,
                        system_context=contexto + _handoff_instrucao_pos_disparo,
                        is_interrupted=lambda: chatLid in self.interrupted_chats,
                    )

                    # Monta resultado com histórico limpo (sem msgs AI pós-interrupção)
                    mensagens_para_contexto = history_for_agent + [{"role": "user", "content": lead_message_atual}]
                    resultado = {
                        "messages": mensagens_para_contexto + [{"role": "assistant", "content": ana_response}]
                    }


                    atualizar_info_lead(chatLid, mensagens_para_contexto, resultado["messages"][-1]["content"])

                    headers = {"client-token": self.zapi_sec_token}

                    if verificar_consulta_disponibilidade(resultado, headers, mensagens_para_contexto):
                        return

                    if verificar_reuniao(resultado, headers, mensagens_para_contexto):
                        return

                    # Fallback: texto parece confirmação de reunião mas não veio como JSON
                    resposta_fallback_text = resultado["messages"][-1]["content"]
                    _lower = resposta_fallback_text.lower()
                    _palavras_conf = [
                        "agendado", "agendada", "marcado", "marcada",
                        "confirmado", "confirmada", "agendei", "marquei", "confirmei",
                    ]
                    _tem_palavra = any(p in _lower for p in _palavras_conf)
                    _tem_horario = bool(re.search(r'(\d{1,2}[h:]\d{0,2})', _lower))
                    _tem_dia = bool(re.search(r'(segunda|terça|terca|quarta|quinta|sexta|sábado|sabado|domingo|\d{1,2}/\d{1,2})', _lower))
                    if _tem_palavra and (_tem_horario or _tem_dia):
                        print(f"{YELLOW}[create_answer] Fallback: texto parece confirmação de reunião. Forçando extração JSON.{RESET}")
                        llm_fb = ChatOpenAI(model="gpt-4o-mini", api_key=os.environ.get("OPENAI_API_KEY"))
                        resp_fb = llm_fb.invoke([
                            SystemMessage(content=get_contexto()),
                            *[
                                HumanMessage(content=m["content"]) if m["role"] == "user"
                                else AIMessage(content=m["content"])
                                for m in mensagens_para_contexto
                            ],
                            AIMessage(content=resposta_fallback_text),
                            SystemMessage(content=(
                                "Sua resposta anterior confirmou uma reunião mas NÃO estava no formato JSON exigido. "
                                "Responda APENAS com o JSON de agendamento, sem nenhum texto antes ou depois:\n"
                                '{"acao": "agendar_reuniao", "data": "YYYY-MM-DD", "horario": "HH:MM", "nome": "Nome Sobrenome", "email": "email@lead.com", "mensagem": "..."}'
                            )),
                        ])
                        resultado_fb = {"messages": resultado["messages"][:-1] + [{"role": "assistant", "content": resp_fb.content}]}
                        if verificar_reuniao(resultado_fb, headers, mensagens_para_contexto):
                            return

                    if verificar_agendamento(resultado, headers):
                        return

                    resposta_completa = resultado["messages"][-1]["content"]

                    # Guard: nunca enviar JSON de ação como mensagem ao lead
                    try:
                        limpa = resposta_completa.strip()
                        if limpa.startswith("```"):
                            limpa = limpa.split("```")[1]
                            if limpa.startswith("json"):
                                limpa = limpa[4:]
                            limpa = limpa.strip()

                        parsed_guard = json.loads(limpa)
                        if isinstance(parsed_guard, dict) and "acao" in parsed_guard:
                            print(f"{RED}Resposta é um JSON de ação não tratado, descartando: {parsed_guard.get('acao')}{RESET}")
                            return
                    except (json.JSONDecodeError, ValueError):
                        pass


                    partes = dividir_em_mensagens(resposta_completa)

                    # Decide citacoes para cada parte (quando lead mandou varias msgs seguidas)
                    citacoes = decidir_citacoes(partes, historico)

                    # Se foi interrompido, cita a mensagem que causou a interrupção na primeira parte
                    if interrupt_reply_id and len(citacoes) > 0:
                        citacoes[0] = interrupt_reply_id
                        print(f" {BLUE} > Citando mensagem interruptora na parte 1: {interrupt_reply_id} {RESET}")

                    for i, parte in enumerate(partes):
                        # Guard: só interrompe se o lead MANDOU nova mensagem (não presença/digitando)
                        if chatLid in self.interrupted_chats:
                            print(f" {YELLOW} > {chatLid} enviou nova mensagem, interrompendo resposta {RESET}")
                            break

                        if i == 0:
                            delay = int(calcular_delay_digitacao(parte, wpm=400))
                        else:
                            time.sleep(0.5)
                            delay = int(calcular_delay_digitacao(parte, wpm=250))
                        print(f" {YELLOW} > Delay digitação: {delay}s para parte {i+1}/{len(partes)} {RESET}")
                        enviar_mensagem(phone, parte, headers, delay_typing=delay, reply_message_id=citacoes[i])
                        print(f"\n\n{GREEN}>Parte {i+1}/{len(partes)} enviada para {chatLid}{RESET}")

                    # Só limpa answer_list se NÃO foi interrompido (senão apaga dados da nova mensagem)
                    if chatLid not in self.interrupted_chats:
                        self.answer_list.pop(chatLid, None)
                    self.agendar_followups(phone, chatLid)

                else:
                    print(f" {RED}\n\n > Chat {chatLid} não permitido {RESET}")
                    return "Desculpe, mas não posso responder a este chat."
            else:
                print(f" {RED} > Nenhuma resposta encontrada para {chatLid} {RESET}")
                return "Desculpe, mas não encontrei uma resposta para esta mensagem."

        # Acumula texto das mensagens recebidas para calcular delay proporcional
        def _extrair_texto_data(d: dict) -> str:
            text = d.get("text", {})
            if isinstance(text, dict):
                return text.get("message", "")
            return str(text) if text else ""

        texto_nova = _extrair_texto_data(data)
        self.pending_texts[chatLid] = self.pending_texts.get(chatLid, "") + " " + texto_nova

        # Delay proporcional: 7s base + 0.04s por caractere acumulado, máx 12s
        total_chars = len(self.pending_texts[chatLid].strip())
        delay = max(7.0, min(7.0 + total_chars * 0.04, 12.0))

        # Debounce: cancela timer anterior do mesmo lead (com lock para evitar race condition)
        with self._timer_lock:
            if chatLid in self.pending_timers:
                self.pending_timers[chatLid].cancel()
                print(f" {YELLOW} > Timer anterior cancelado para {chatLid} {RESET}")

            print(f" {YELLOW} > Delay proporcional: {delay:.1f}s ({total_chars} chars) para {chatLid} {RESET}")
            timer = threading.Timer(delay, create_answer, args=(phone, chatLid, data))
            self.pending_timers[chatLid] = timer
            timer.start()
